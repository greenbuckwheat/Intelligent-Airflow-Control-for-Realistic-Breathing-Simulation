"""English edition of the current, plain-text PPV meeting slides."""
import re
from pathlib import Path

import build_ppv_group_meeting_ppt as deck
from matplotlib.text import Text
from pptx.util import Pt, Inches


TRANSLATIONS = {
    "理想值": "Ideal",
    "up1 升扫": "Scan 1: upward",
    "scan2 升扫": "Scan 2: upward",
    "scan2 降扫": "Scan 2: downward",
    "实测电压 / V": "Measured voltage / V",
    "误差 / mV": "Error / mV",
    "理想目标电压 / V": "Ideal target voltage / V",
    "目标电压 / V": "Target voltage / V",
    "升扫 − 降扫": "Upward minus downward",
    "同码值电压差 / mV": "Voltage difference at the same code / mV",
    "第二次升扫 − 第一次升扫": "Second minus first upward scan",
    "4.2 V 夹闭": "4.2 V: closed",
    "ATM 打开后时间 / s": "Time after ATM opening / s",
    "零点校正流量 / SLM": "Baseline-corrected flow / SLM",
    "单次实验": "Individual trial",
    "各电压中位数": "Median at each voltage",
    "分段物理模型": "Piecewise fitted model",
    "低电压平台": "Low-voltage plateau",
    "PPV 命令电压 / V": "PPV command voltage / V",
    "有效吸气时间常数 / s": "Effective inspiratory time constant / s",
    "原假设 $a=1-U/5$": "Original assumption $a=1-U/5$",
    "实测分段 $a_{eff}$": "Fitted piecewise $a_{eff}$",
    "归一化有效开度": "Normalized effective opening",
    "10% 参考线": "10% reference",
    "时间常数变异系数 $CV$ / %": "Time-constant coefficient of variation $CV$ / %",
    "PPV 电压 / V": "PPV voltage / V",
    "PPV40V 实物实验结果": "PPV40V Bench Experiments",
    "PWM—电压校准与 fake lung 吸气开度辨识": "PWM–voltage calibration and fake-lung inspiration modelling",
    "两组实验": "Two experiments",
    "1. PCB/Arduino PWM → 电压": "1. PCB/Arduino PWM → voltage",
    "2. PPV 电压 → 吸气流量": "2. PPV voltage → inspiratory flow",
    "两组实验分别回答什么？": "What do the two experiments establish?",
    "先区分电气命令链和气动被控对象": "Electrical command conversion and pneumatic system response",
    "实验 1：PWM—电压扫描": "Experiment 1: PWM–voltage sweep",
    "滤波/电平转换": "Filter / scaling",
    "万用表": "Multimeter",
    "线性、稳态误差、升降扫重复性": "Linearity, steady-state error and sweep repeatability",
    "不测阀位、流量或动态延迟": "Valve position, flow and dynamic delay were not measured",
    "实验 2：PPV—fake lung 吸气": "Experiment 2: fake-lung inspiration",
    "PPV 电压": "PPV voltage",
    "ATM 开启": "Open ATM",
    "吸气阶跃": "Inspiration step",
    "50 Hz 流量": "Flow at 50 Hz",
    "一阶时间常数、有效调节区、夹闭点、γ": "Time constant, control range, closure voltage and γ",
    "PPV 已提前设定，不能隔离 PPV 动作延迟": "PPV was preset; its actuation delay cannot be isolated",
    "联合分析才能判断：流量死区不是电压输出失真造成的。": "The electrical calibration supports interpreting the flow plateau as a pneumatic-system effect.",
    "实验 1：经过 PCB/Arduino 后，PWM 与输出电压基本一致": "Experiment 1: output voltage closely follows the PWM command",
    "D5 1 kHz、12-bit PWM → PCB 滤波/电平转换 → PPV 0–5 V 控制输入": "D5: 1 kHz, 12-bit PWM → PCB filtering / scaling → PPV 0–5 V control input",
    "实验方法": "Method",
    "0–100% 升扫，再 100–0% 降扫": "Sweep 0–100%, then 100–0%",
    "数据结果": "Measured results",
    "• 线性 R²：0.999993\n• RMSE：4.63 mV\n• 最大误差：7.4 mV\n• 升/降扫最大差：1.0 mV": "• Linear R²: 0.999993\n• RMSE: 4.63 mV\n• Maximum error: 7.4 mV\n• Max. up/down difference: 1.0 mV",
    "实验结论\n最大误差仅 0.148% FS，线性换算已足够。": "Conclusion\nMaximum error: 0.148% FS. A linear conversion is adequate for initial modelling.",
    "因此第二个实验中的流量非线性不是 PCB/Arduino 电压输出偏差造成的。": "The small electrical error does not explain the substantial flow nonlinearity in Experiment 2.",
    "实验 2：固定 PPV 电压下的 fake lung 吸气": "Experiment 2: inspiration at a fixed PPV voltage",
    "每次手动压扁 fake lung，再用 ATM 开启作为吸气阶跃": "Manually compress the fake lung, then open ATM to initiate inspiration",
    "压扁 fake lung": "Compress lung",
    "手动初始状态": "Manual initial state",
    "ATM/VAC 关闭": "Close ATM/VAC",
    "隔离气路": "Isolate the circuit",
    "设置 PPV": "Set PPV",
    "打开 ATM": "Open ATM",
    "正常吸气": "Inspiration",
    "记录流量": "Record flow",
    "40 份 CSV": "40 CSV files",
    "包含不同电压和重复试验": "Multiple voltages and repeated trials",
    "38 次动态拟合": "38 dynamic fits",
    "37.csv 剔除；42.csv 作为夹闭点": "Exclude 37.csv; retain 42.csv as the closed point",
    "R² 中位数 0.9994": "Median R²: 0.9994",
    "累积体积的一阶拟合": "First-order cumulative-volume fits",
    "未记录压力": "No pressure data",
    "不能分离阀阻与 fake lung 顺应性": "Valve resistance and lung compliance cannot be separated",
    "可识别整个“PPV—管路—fake lung”的有效动态，不是 PPV 独立流量系数。": "The fits characterize the combined PPV–tubing–fake-lung dynamics.",
    "a. 初步建模参数如何从流量数据得到？": "a. From flow measurements to initial model parameters",
    "每次试验给出一条 Q(t)，同一电压的重复试验再取稳健汇总": "Fit each Q(t) trace, then summarize repeated trials at each voltage using medians",
    "零点校正": "Zero correction",
    "ATM 前 2 s 中位数": "Median before ATM\n(last 2 s)",
    "累积吸入体积": "Inspired volume",
    "一阶拟合": "First-order fit",
    "分电压汇总": "Group by voltage",
    "τeff 中位数": "Median τeff",
    "PPV 模型": "PPV model",
    "单次实验可得到": "From each trial",
    "• 峰值/早期流量、累积体积 V(t)\n• τeff、拟合体积 V∞、综合延迟 td": "• Peak / early flow and cumulative volume V(t)\n• τeff, fitted V∞ and combined response delay td",
    "多电压数据可进一步得到": "From the voltage sweep",
    "• 全开基础时间常数 τin\n• 有效调节起点 Ud、夹闭点 Uc、开度指数 γ": "• Fully open baseline time constant τin\n• Control onset Ud, closure Uc and opening exponent γ",
    "峰值流量对初始压扁程度很敏感；因此 γ 主要根据一阶时间常数辨识。": "Peak flow depends strongly on initial compression; γ is therefore estimated mainly from time constants.",
    "代表性流量曲线：3.3 V 后逐渐变慢": "Representative flow traces: slower inspiration above 3.3 V",
    "每个电压选取最接近该组中位时间常数的一次实验": "For each voltage, show the trial with a time constant closest to the group median",
    "0–3.2 V：衰减时间接近，阀限制较弱": "0–3.2 V: similar decay times; weak valve restriction",
    "3.5–3.9 V：峰值下降，吸气持续时间增加": "3.5–3.9 V: lower peaks and longer inspiration",
    "4.1 V：峰值约 5.6 SLM，响应显著变慢": "4.1 V: peak ≈5.6 SLM; much slower response",
    "4.2 V：流量与零点噪声相当，确认夹闭": "4.2 V: flow near baseline noise; closure confirmed",
    "峰值流量受手动压扁程度影响，不用于单独辨识 γ": "Initial compression affects peak flow; do not use peaks alone to identify γ",
    "曲线展示了明确的整体趋势，但局部非单调主要来自初始机械状态差异。": "A clear overall trend is present; differences in initial mechanical state likely explain local non-monotonicity.",
    "一阶时间常数辨识": "First-order time-constant identification",
    "对零点校正流量积分，拟合累积吸入体积": "Integrate baseline-corrected flow and fit cumulative inspired volume",
    "拟合模型": "Fitted model",
    "初步结果": "Initial results",
    "• 0–3.2 V：τ 中位数 1.91 s\n• 4.1 V：τ = 9.41 s\n• 3.9 V 以上增长很快": "• 0–3.2 V: median τ = 1.91 s\n• 4.1 V: τ = 9.41 s\n• Rapid increase above 3.9 V",
    "全开基础时间常数可取 τin ≈ 1.9 s；3.9 V 以上时间常数快速增大。": "Use τin ≈ 1.9 s as the fully open baseline; the time constant rises rapidly above 3.9 V.",
    "PPV 非线性：需要分段有效开度": "PPV nonlinearity: a piecewise effective opening",
    "低电压死区 + 4.2 V 提前夹闭，不适合全程 a=1−U/5": "A low-voltage plateau and closure at 4.2 V make a=1−U/5 unsuitable across the full range",
    "分段模型拟合结果": "Piecewise model fit",
    "• 平台终点 Ud = ": "• Plateau endpoint Ud = ",
    "• 有效开度 γ = ": "• Opening exponent γ = ",
    "• 时间常数拟合 R² = ": "• Time-constant fit R² = ",
    "重复选择敏感范围：γ = 0.60–0.68": "Trial-selection sensitivity: γ = 0.60–0.68",
    "建议用 γ≈0.65 + 分段 aeff；不建议将局部拟合的 γ≈2.3–2.6 当作全局物理参数。": "Use γ≈0.65 with piecewise aeff; the earlier local γ≈2.3–2.6 is not a global physical parameter.",
    "b. 数据筛选与优化": "b. Data screening and improvements",
    "不把所有偏差都当作阀门特性：先分离记录错误、夹闭点和低重复性试验": "Account for recording errors, confirmed closure and trial-to-trial variability",
    "文件名 3.7 V；遥测 PPV 全程 0 V。": "Filename: 3.7 V; recorded PPV: 0 V throughout.",
    "反复 ATM_ON 仍无流量；4.2 V 确认夹闭。": "No flow after repeated ATM_ON; 4.2 V closure confirmed.",
    "与 33.1 差异大；后段有反向波动。": "Differs from 33.1.csv; reverse-flow fluctuations late in the trial.",
    "记录后才设 3.5 V；结果偏低。": "3.5 V set after recording began; lower-flow results.",
    "剔除": "Exclude",
    "保留闭点": "Retain closure",
    "低信心度": "Low confidence",
    "35.2/35.3 在 ATM 开启前已稳定约 3.26/3.77 s，不因操作顺序直接删除。": "35.2/35.3 settled for 3.26/3.77 s before ATM; retain with caution.",
    "当前最大的实验误差源是每次 fake lung 初始压扁程度和机械状态不一致。": "The main uncertainty is inconsistent initial compression and mechanical state of the fake lung.",
    "c. 误差分析：本次只辨识了吸气过程": "c. Limitations: only inspiration was identified",
    "当前的 τ 是 ATM 吸气路径的 τin/τeff，并没有得到真空呼气路径的 τout": "The fitted τ describes the ATM inspiratory path; the vacuum expiratory time constant τout is still unknown",
    "已得到": "Identified",
    "吸气流量 Q(t) 和积分体积 V(t)": "Inspiratory flow Q(t) and integrated volume V(t)",
    "全开 τin 与部分开度 τeff": "Fully open τin and restricted-opening τeff",
    "PPV 有效调节区、夹闭点和 γ": "PPV control range, closure voltage and γ",
    "未得到": "Still unknown",
    "呼气/真空收缩时间常数 τout": "Expiratory / vacuum-collapse time constant τout",
    "PPV 命令到流量的独立纯延迟": "Isolated PPV command-to-flow delay",
    "阀前后压差：不能分离阀阻与顺应性": "Pressure drop: resistance and compliance cannot be separated",
    "工程影响": "Control implications",
    "真空源可调，τout 对当前吸气建模优先级较低": "Vacuum is adjustable; τout is less urgent for the current inspiration model",
    "但要准确评估呼吸频率、I:E 和循环一致性，仍必须测 τout": "τout is still needed to assess breathing rate, I:E ratio and cycle consistency",
    "建议在多个 vacuum 设定下分别辨识": "Identify τout at several vacuum settings",
    "可调 vacuum 可用来修正呼气速度，但不能替代 τout 标定；完整控制模型仍需要吸气和呼气两组参数。": "Adjustable vacuum can tune expiration speed; a complete model still needs both inspiratory and expiratory parameters.",
    "初步实验结果：数据不完美，但足以给出可用的初值": "Preliminary results: usable initial model parameters",
    "参数用于初步模型和敏感性分析，不应当作最终高精度标定": "Use these estimates for initial modelling and sensitivity analysis, rather than final calibration",
    "全开吸气；1.8–2.1 s": "Fully open; 1.8–2.1 s",
    "开始明显调节": "Control onset",
    "确认完全夹闭": "Confirmed closure",
    "敏感范围 0.60–0.70": "Sensitivity: 0.60–0.70",
    "τ—电压拟合": "τ–voltage fit",
    "建议的吸气初步模型": "Suggested initial inspiration model",
    "aeff=1（U≤3.1 V）；3.1–4.2 V 线性降至 0；U≥4.2 V 夹闭。": "aeff=1 for U≤3.1 V; decreases linearly to 0 at 4.2 V; closed for U≥4.2 V.",
    "该结果可以支持": "Applications",
    "初步前馈模型参数化": "Initial feedforward model parameters",
    "控制器中的参数敏感性测试": "Controller parameter sensitivity tests",
    "后续硬件循环实验的基准组": "Baseline for subsequent breathing-cycle tests",
    "关键是使用“分段有效开度 + γ≈0.65”，不是保留 a=1−U/5 后只修改 γ。": "Apply γ≈0.65 together with the piecewise effective opening aeff(U).",
    "Future work：今天下午的测试计划": "Future work: this afternoon's test plan",
    "今天已拿到 vacuum source；下午将补齐呼气参数，并观察整个控制在多个呼吸循环中的表现": "Vacuum source received today: measure expiration and evaluate control over consecutive breathing cycles",
    "Vacuum/呼气实验": "Vacuum / expiration tests",
    "在多个负压设定下记录流量，辨识 τout 和收缩体积。": "Record flow at several vacuum settings; estimate τout and collapse volume.",
    "PPV 阶跃与延迟": "PPV step response and delay",
    "ATM 已打开、压差稳定时单独阶跃 PPV，同步记录命令和流量。": "With ATM open and pressure drop stable, step PPV and record command and flow together.",
    "整个呼吸循环": "Complete breathing cycles",
    "运行多个连续吸气/呼气循环，检查相位切换和循环间重复性。": "Run consecutive inspiration / expiration cycles; check phase switching and repeatability.",
    "控制性能记录": "Control performance logging",
    "同步保存目标/实测流量、积分体积、PPV 命令及 ATM/VAC 状态。": "Log target / measured flow, integrated volume, PPV command and ATM/VAC states together.",
    "下午的成功标志：获得 τout，并在多个呼吸循环中同时看到稳定的流量跟踪、体积和阀门切换。": "Success criteria: estimate τout and demonstrate consistent flow tracking, volume and valve switching over multiple cycles.",
    "附录：各电压的时间常数汇总": "Appendix: time constants by voltage",
    "表内范围为同一电压的最小—最大单次拟合值": "Ranges show the minimum and maximum individual fitted values at each voltage",
    "电压/V": "Voltage / V",
    "τ中位/s": "Median τ / s",
    "τ范围/s": "τ range / s",
    "峰值流量/SLM": "Peak flow / SLM",
    "体积中位/mL": "Median volume / mL",
    "夹闭": "Closed",
}


def translate(value):
    if value in TRANSLATIONS:
        return TRANSLATIONS[value]
    for source, target in TRANSLATIONS.items():
        if source.startswith("• ") and source.endswith(" = "):
            value = value.replace(source, target)
    if re.search(r"[\u4e00-\u9fff]", value):
        raise ValueError(f"Missing translation: {value}")
    return value


original_add_text = deck.add_text
original_add_bullets = deck.add_bullets
original_plot_text = Text.set_text


def english_text(slide, x, y, w, h, text, size=16, *args, **kwargs):
    translated = translate(text)
    size = min(size, 25) if y < 0.6 else size * 0.93
    # English lines need wrapping; allow slightly more height within the same layout.
    box = original_add_text(slide, x, y, w, max(h, 0.38), translated, size, *args, **kwargs)
    box.text_frame.word_wrap = True
    for p in box.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Arial"
    if h < 0.6 and y > 1.3 and y < 6.6:
        box.height = Inches(max(h, 0.52))
    return box


def english_bullets(slide, x, y, w, h, items, size=16, **kwargs):
    box = original_add_bullets(slide, x, y, w, h, [translate(s) for s in items], size * 0.90, **kwargs)
    for p in box.text_frame.paragraphs:
        p.font.name = "Arial"
    return box


def english_plot_text(self, text):
    return original_plot_text(self, translate(str(text)) if text is not None else text)


def main():
    deck.ASSETS = deck.ANALYSIS / "ppt_assets_en"
    deck.OUTPUT = deck.EXPERIMENT / "PPV40V_Two_Experiments_Group_Meeting_2026-10-06_EN.pptx"
    deck.plt.rcParams["font.family"] = "Arial"
    deck.add_text = english_text
    deck.add_bullets = english_bullets
    Text.set_text = english_plot_text
    _, _, _, summary, model = deck.prepare_assets()
    deck.build_presentation(summary, model)
    prs = deck.Presentation(deck.OUTPUT)
    prs.core_properties.title = "PPV40V: Electrical Calibration and Inspiratory Dynamics"
    prs.core_properties.subject = "PWM–voltage sweep and fake-lung inspiration experiments"
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.text = translate(run.text)
                                run.font.name = "Arial"
                                run.font.size = Pt(10)
    prs.save(deck.OUTPUT)
    print(deck.OUTPUT)


if __name__ == "__main__":
    main()
