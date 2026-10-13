"""Insert the existing digital-twin result into the 2026-10-02 poster.

Only the Experimental results placeholder and its caption are changed. The
supplied PPTX/HTML remain untouched; new artifacts are written beside this file.
"""

from __future__ import annotations

import base64
import re
import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


POSTER_DIR = Path(__file__).resolve().parent
SOURCE_DIR = Path(r"D:\Downloads")
SOURCE_PPTX = SOURCE_DIR / "LungLogic_Poster_2026-10-02.pptx"
SOURCE_HTML = SOURCE_DIR / "LungLogic_Poster_2026-10-02.html"
RESULT_IMAGE = POSTER_DIR / "assets" / "single_lung_validation_two_panel.png"
STEM = "LungLogic_Poster_2026-10-02_with_simulation"
OUTPUT_PPTX = POSTER_DIR / f"{STEM}.pptx"
OUTPUT_PDF = POSTER_DIR / f"{STEM}.pdf"
OUTPUT_HTML = POSTER_DIR / f"{STEM}.html"
OUTPUT_PREVIEW = POSTER_DIR / f"{STEM}_preview.png"

OLD_CAPTION = "Fig. 5. Reserved for measured hardware results."
NEW_CAPTION = "Fig. 5. Digital-twin simulation at 15 bpm."


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def image_fit(box) -> tuple[int, int, int, int]:
    left, top, width, height = box
    with Image.open(RESULT_IMAGE) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        actual_width = width
        actual_height = round(width / image_ratio)
    else:
        actual_height = height
        actual_width = round(height * image_ratio)
    return (
        left + (width - actual_width) // 2,
        top + (height - actual_height) // 2,
        actual_width,
        actual_height,
    )


def insert_in_pptx() -> tuple[int, int, int, int]:
    shutil.copy2(SOURCE_PPTX, OUTPUT_PPTX)
    deck = Presentation(OUTPUT_PPTX)
    if len(deck.slides) != 1:
        raise ValueError("Expected a one-page poster")
    slide = deck.slides[0]
    placeholder = [s for s in slide.shapes if s.has_text_frame and
                   s.text.strip() == "EXPERIMENTAL RESULTS TO BE ADDED"]
    note = [s for s in slide.shapes if s.has_text_frame and
            s.text.startswith("Insert measured flow and volume plots")]
    caption = [s for s in slide.shapes if s.has_text_frame and
               s.text.strip() == OLD_CAPTION]
    frames = [s for s in slide.shapes if s.has_text_frame and not s.text.strip() and
              abs(s.left / Inches(1) - 0.63) < 0.01 and
              abs(s.top / Inches(1) - 38.14) < 0.01 and
              abs(s.width / Inches(1) - 20.55) < 0.01]
    if not all(len(group) == 1 for group in (placeholder, note, caption, frames)):
        raise ValueError("Experimental-results placeholder did not match the supplied deck")

    box = (frames[0].left, frames[0].top, frames[0].width, frames[0].height)
    for shape in (frames[0], placeholder[0], note[0]):
        remove_shape(shape)
    # Preserve all caption typography and positioning; only the text changes.
    cap = caption[0]
    cap.text_frame.paragraphs[0].runs[0].text = NEW_CAPTION
    fitted = image_fit(box)
    slide.shapes.add_picture(str(RESULT_IMAGE), *fitted)
    deck.save(OUTPUT_PPTX)
    return fitted


def _one_sub(pattern: str, replacement: str, document: str) -> str:
    updated, count = re.subn(pattern, replacement, document, count=1, flags=re.DOTALL)
    if count != 1:
        raise ValueError(f"HTML placeholder pattern not found: {pattern[:70]}")
    return updated


def insert_in_html(fitted: tuple[int, int, int, int]) -> None:
    document = SOURCE_HTML.read_text(encoding="utf-8")
    document = _one_sub(
        r'<div class="item shape" style="left:1\.6002cm;top:96\.8756cm;'
        r'width:52\.1970cm;height:14\.3764cm;[^\"]*"></div>',
        "", document,
    )
    document = _one_sub(
        r'<div class="item text" style="[^\"]*">'
        r'EXPERIMENTAL RESULTS TO BE ADDED</div>', "", document,
    )
    document = _one_sub(
        r'<div class="item text" style="[^\"]*">'
        r'Insert measured flow and volume plots<br>Report: VT error.*?</div>',
        "", document,
    )
    if document.count(OLD_CAPTION) != 1:
        raise ValueError("Expected one HTML results caption")
    document = document.replace(OLD_CAPTION, NEW_CAPTION, 1)

    left, top, width, height = fitted
    to_cm = lambda emu: emu / 360000
    image_style = (
        f"left:{to_cm(left):.4f}cm;top:{to_cm(top):.4f}cm;"
        f"width:{to_cm(width):.4f}cm;height:{to_cm(height):.4f}cm;"
    )
    data = base64.b64encode(RESULT_IMAGE.read_bytes()).decode("ascii")
    image_html = (
        f'<img class="item picture" style="{image_style}" '
        f'src="data:image/png;base64,{data}" '
        'alt="Digital-twin simulated volume and flow tracking">'
    )
    caption_marker = re.compile(
        r'(?=<div class="item text" style="[^\"]*">'
        r'Fig\. 5\. Digital-twin simulation at 15 bpm)'
    )
    document, count = caption_marker.subn(lambda _: image_html, document, count=1)
    if count != 1:
        raise ValueError("Could not place result image before HTML caption")
    OUTPUT_HTML.write_text(document, encoding="utf-8")


def export_pdf_and_preview() -> None:
    import win32com.client

    powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
    powerpoint.Visible = True
    try:
        deck = powerpoint.Presentations.Open(str(OUTPUT_PPTX.resolve()), WithWindow=False)
        deck.SaveAs(str(OUTPUT_PDF.resolve()), 32)
        deck.Slides(1).Export(str(OUTPUT_PREVIEW.resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        powerpoint.Quit()


def main() -> None:
    for path in (SOURCE_PPTX, SOURCE_HTML, RESULT_IMAGE):
        if not path.exists():
            raise FileNotFoundError(path)
    fitted = insert_in_pptx()
    insert_in_html(fitted)
    export_pdf_and_preview()
    for path in (OUTPUT_PPTX, OUTPUT_PDF, OUTPUT_HTML, OUTPUT_PREVIEW):
        print(path)


if __name__ == "__main__":
    main()
