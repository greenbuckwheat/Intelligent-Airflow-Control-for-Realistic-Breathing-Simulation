"""Insert the supplied single-branch prototype and PCB/STM32 photographs."""

from __future__ import annotations

import base64
from pathlib import Path

from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


HERE = Path(__file__).resolve().parent
SOURCE = "LungLogic_Poster_2026-10-02_single_branch_compact"
OUTPUT = "LungLogic_Poster_2026-10-02_prototype_pcb"
PROTOTYPE = HERE / "2.jpg"
PCB = HERE / "1.jpg"


def fit_image(path: Path, left: float, top: float, width: float, height: float):
    with Image.open(path) as image:
        aspect = image.width / image.height
    if width / height > aspect:
        fit_height = height
        fit_width = height * aspect
    else:
        fit_width = width
        fit_height = width / aspect
    return (
        left + (width - fit_width) / 2,
        top + (height - fit_height) / 2,
        fit_width,
        fit_height,
    )


def remove(shape):
    element = shape._element
    element.getparent().remove(element)


def edit_pptx():
    deck = Presentation(HERE / f"{SOURCE}.pptx")
    slide = deck.slides[0]
    by_name = {shape.name: shape for shape in slide.shapes}
    for name in ("Rounded Rectangle 25", "TextBox 26", "Picture 41"):
        remove(by_name[name])

    prototype_box = (24.58, 9.25, 7.54, 7.45)
    pcb_box = (28.40, 19.84, 3.66, 3.15)
    for path, box in ((PROTOTYPE, prototype_box), (PCB, pcb_box)):
        left, top, width, height = fit_image(path, *box)
        slide.shapes.add_picture(
            str(path), Inches(left), Inches(top), Inches(width), Inches(height)
        )

    captions = {
        "Fig. 1. Single-lung route; prototype photograph pending.":
            "Fig. 1. Single-lung pneumatic route and 250 mL prototype.",
        "Fig. 2. Control architecture and bridge PCB.":
            "Fig. 2. Control architecture and bridge PCB connected to STM32.",
    }
    changed = set()
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        original = shape.text.strip()
        if original in captions:
            shape.text_frame.paragraphs[0].runs[0].text = captions[original]
            changed.add(original)
    if changed != set(captions):
        raise ValueError(f"Missing captions: {set(captions) - changed}")
    deck.save(HERE / f"{OUTPUT}.pptx")


def html_image(soup, path: Path, box):
    left, top, width, height = fit_image(path, *box)
    node = soup.new_tag("img")
    node["class"] = ["item", "picture"]
    node["style"] = (
        f"left:{left * 2.54:.4f}cm;top:{top * 2.54:.4f}cm;"
        f"width:{width * 2.54:.4f}cm;height:{height * 2.54:.4f}cm;"
    )
    node["src"] = "data:image/jpeg;base64," + base64.b64encode(path.read_bytes()).decode("ascii")
    node["alt"] = "Single-branch prototype" if path == PROTOTYPE else "Bridge PCB connected to STM32"
    return node


def edit_html():
    soup = BeautifulSoup((HERE / f"{SOURCE}.html").read_text(encoding="utf-8"), "html.parser")
    items = soup.select(".item")
    placeholder = next(item for item in items if "Prototype photograph" in item.get_text())
    frame = next(item for item in items if item.get("style", "").startswith("left:62.4332cm;top:23.4950cm"))
    pcb = next(item for item in items if item.name == "img" and item.get("style", "").startswith("left:73.0075cm;top:50.4952cm"))
    frame.replace_with(html_image(soup, PROTOTYPE, (24.58, 9.25, 7.54, 7.45)))
    placeholder.decompose()
    pcb.replace_with(html_image(soup, PCB, (28.40, 19.84, 3.66, 3.15)))
    captions = {
        "Fig. 1. Single-lung route; prototype photograph pending.":
            "Fig. 1. Single-lung pneumatic route and 250 mL prototype.",
        "Fig. 2. Control architecture and bridge PCB.":
            "Fig. 2. Control architecture and bridge PCB connected to STM32.",
    }
    for item in soup.select(".item.text"):
        original = item.get_text(strip=True)
        if original in captions:
            item.string = captions[original]
    (HERE / f"{OUTPUT}.html").write_text(str(soup), encoding="utf-8")


def export():
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{OUTPUT}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{OUTPUT}.pdf").resolve()), 32)
        deck.Slides(1).Export(str((HERE / f"{OUTPUT}_preview.png").resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    edit_pptx()
    edit_html()
    export()
    print(HERE / f"{OUTPUT}.pdf")
