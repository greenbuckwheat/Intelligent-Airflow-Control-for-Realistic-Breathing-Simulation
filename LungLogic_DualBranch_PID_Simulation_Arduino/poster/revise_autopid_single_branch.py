"""Clarify poster control text and keep future work scoped to one 250 mL branch."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


POSTER_DIR = Path(__file__).resolve().parent
SOURCE_STEM = "LungLogic_Poster_2026-10-02_with_simulation_no_formula"
OUTPUT_STEM = "LungLogic_Poster_2026-10-02_single_branch_compact"
SOURCE_PPTX = POSTER_DIR / f"{SOURCE_STEM}.pptx"
SOURCE_HTML = POSTER_DIR / f"{SOURCE_STEM}.html"
OUTPUT_PPTX = POSTER_DIR / f"{OUTPUT_STEM}.pptx"
OUTPUT_PDF = POSTER_DIR / f"{OUTPUT_STEM}.pdf"
OUTPUT_HTML = POSTER_DIR / f"{OUTPUT_STEM}.html"
OUTPUT_PREVIEW = POSTER_DIR / f"{OUTPUT_STEM}_preview.png"

REVISIONS = {
    "3.4 Auto PID and measured-delay compensation":
        "3.4 Optional Auto PID and measured-delay compensation",
    "Auto PID scores bounded Kp/Ki trials. Smith prediction uses the identified model and measured delay; hardware adds no artificial delay.":
        "Auto PID scores bounded Kp/Ki over full runs by flow, volume and saturation; optional Smith prediction uses the model and measured delay.",
    "Integrate the biomimetic lung model":
        "Integrate the 250 mL branch",
    "Re-identify the model and delay":
        "Re-identify model and delay",
    "Validate bilateral repeatability":
        "Validate branch repeatability",
    "Extend the verified branch to independent left/right control across the required breathing range.":
        "Connect the single branch to the biomimetic lung model and test the required breathing range.",
}


def edit_pptx() -> None:
    shutil.copy2(SOURCE_PPTX, OUTPUT_PPTX)
    deck = Presentation(OUTPUT_PPTX)
    if len(deck.slides) != 1:
        raise ValueError("Expected a one-page poster")
    slide = deck.slides[0]
    changed = set()
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        old_text = shape.text.strip()
        if old_text not in REVISIONS:
            continue
        if old_text in changed:
            raise ValueError(f"Duplicate text box: {old_text}")
        paragraph = shape.text_frame.paragraphs[0]
        if len(paragraph.runs) != 1:
            raise ValueError(f"Unexpected rich text: {old_text}")
        paragraph.runs[0].text = REVISIONS[old_text]
        if old_text.startswith("Auto PID scores"):
            shape.top = Inches(34.82)
            shape.height = Inches(0.76)
        changed.add(old_text)
    if changed != set(REVISIONS):
        raise ValueError(f"Missing expected text: {set(REVISIONS) - changed}")
    deck.save(OUTPUT_PPTX)


def edit_html() -> None:
    document = SOURCE_HTML.read_text(encoding="utf-8")
    for old, new in REVISIONS.items():
        if document.count(old) != 1:
            raise ValueError(f"Expected text exactly once in HTML: {old}")
        document = document.replace(old, new, 1)
    # Match the small text box move made in the PPTX to hold the two-line copy.
    pattern = (
        r'(<div class="item text" style="left:1\.5748cm;top:)88\.5698cm'
        r'(;width:80\.4164cm;height:)2\.4384cm'
        r'([^\"]*">Auto PID scores bounded Kp/Ki)'
    )
    document, count = re.subn(
        pattern, r'\g<1>88.4428cm\g<2>1.9304cm\g<3>', document, count=1,
    )
    if count != 1:
        raise ValueError("Could not update HTML paragraph dimensions")
    OUTPUT_HTML.write_text(document, encoding="utf-8")


def export_pdf_and_preview() -> None:
    import win32com.client

    powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
    powerpoint.Visible = True
    try:
        deck = powerpoint.Presentations.Open(str(OUTPUT_PPTX.resolve()), WithWindow=False)
        deck.SaveAs(str(OUTPUT_PDF.resolve()), 32)
        deck.Slides(1).Export(str(OUTPUT_PREVIEW.resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        powerpoint.Quit()


def main() -> None:
    for path in (SOURCE_PPTX, SOURCE_HTML):
        if not path.exists():
            raise FileNotFoundError(path)
    edit_pptx()
    edit_html()
    export_pdf_and_preview()
    for path in (OUTPUT_PPTX, OUTPUT_PDF, OUTPUT_HTML, OUTPUT_PREVIEW):
        print(path)


if __name__ == "__main__":
    main()
