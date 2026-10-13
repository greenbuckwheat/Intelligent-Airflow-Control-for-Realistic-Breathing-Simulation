"""Focus the poster's future work on physical biomimetic-lung integration."""

from pathlib import Path

from pptx import Presentation


HERE = Path(__file__).resolve().parent
SOURCE = "LungLogic_Poster_2026-10-02_prototype_pcb"
OUTPUT = "LungLogic_Poster_2026-10-02_physical_lung_future_work"

REVISIONS = {
    "Integrate the 250 mL branch": "Connect to the physical lung model",
    "Re-identify model and delay": "Identify coupled dynamics and delay",
    "Validate branch repeatability": "Validate physical-model breathing",
    "Connect the single branch to the biomimetic lung model and test the required breathing range.":
        "Connect the 250 mL control branch to the physical biomimetic lung model, then test tidal-volume tracking and repeatability.",
}


def edit_pptx():
    deck = Presentation(HERE / f"{SOURCE}.pptx")
    changed = set()
    for shape in deck.slides[0].shapes:
        if not shape.has_text_frame:
            continue
        old = shape.text.strip()
        if old in REVISIONS:
            paragraph = shape.text_frame.paragraphs[0]
            if len(paragraph.runs) != 1:
                raise ValueError(f"Unexpected rich text: {old}")
            paragraph.runs[0].text = REVISIONS[old]
            changed.add(old)
    if changed != set(REVISIONS):
        raise ValueError(f"Missing text: {set(REVISIONS) - changed}")
    deck.save(HERE / f"{OUTPUT}.pptx")


def edit_html():
    document = (HERE / f"{SOURCE}.html").read_text(encoding="utf-8")
    for old, new in REVISIONS.items():
        if document.count(old) != 1:
            raise ValueError(f"Expected one HTML occurrence: {old}")
        document = document.replace(old, new, 1)
    (HERE / f"{OUTPUT}.html").write_text(document, encoding="utf-8")


def export():
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{OUTPUT}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{OUTPUT}.pdf").resolve()), 32)
        deck.Slides(1).Export(str((HERE / f"{OUTPUT}_preview.png").resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    edit_pptx()
    edit_html()
    export()
    print(HERE / f"{OUTPUT}.pdf")
