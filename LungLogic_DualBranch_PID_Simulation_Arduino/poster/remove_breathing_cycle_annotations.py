"""Replace the breathing-cycle plot with the clean, unannotated version."""

from __future__ import annotations

import base64
import re
from pathlib import Path

from pptx import Presentation


HERE = Path(__file__).resolve().parent
IMAGE = HERE / "assets" / "nominal_breathing_cycle.png"
STEMS = (
    "LungLogic_Poster_2026-10-02_prototype_pcb",
    "LungLogic_Poster_2026-10-02_physical_lung_future_work",
)
BOX_CM = (7.3834, 36.2, 27.2955, 9.017)


def update_pptx(stem: str) -> None:
    path = HERE / f"{stem}.pptx"
    deck = Presentation(path)
    slide = deck.slides[0]
    left, top, width, height = BOX_CM
    pictures = [
        shape for shape in slide.shapes
        if shape.shape_type == 13
        and abs(shape.left / 360000 - left) < 0.01
        and abs(shape.top / 360000 - top) < 0.01
        and abs(shape.width / 360000 - width) < 0.01
        and abs(shape.height / 360000 - height) < 0.01
    ]
    if len(pictures) != 1:
        raise ValueError(f"Breathing-cycle picture not found in {stem}")
    pictures[0]._element.getparent().remove(pictures[0]._element)
    slide.shapes.add_picture(
        str(IMAGE),
        round(left * 360000),
        round(top * 360000),
        round(width * 360000),
        round(height * 360000),
    )
    deck.save(path)


def update_html(stem: str) -> None:
    path = HERE / f"{stem}.html"
    document = path.read_text(encoding="utf-8")
    encoded = base64.b64encode(IMAGE.read_bytes()).decode("ascii")
    replacement = (
        '<img alt="Nominal single-branch breathing cycle" class="item picture" '
        f'src="data:image/png;base64,{encoded}" '
        'style="left:7.3834cm;top:36.2000cm;width:27.2955cm;height:9.0170cm;"/>'
    )
    pattern = re.compile(
        r'<img [^>]*?src="data:image/png;base64,[^"]+" '
        r'style="left:7\.3834cm;top:36\.2000cm;width:27\.2955cm;height:9\.0170cm;"/>'
    )
    document, count = pattern.subn(lambda _: replacement, document, count=1)
    if count != 1:
        raise ValueError(f"Breathing-cycle image not found in HTML: {stem}")
    path.write_text(document, encoding="utf-8")


def export(stem: str) -> None:
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{stem}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{stem}.pdf").resolve()), 32)
        deck.Slides(1).Export(
            str((HERE / f"{stem}_preview.png").resolve()), "PNG", 2384, 3370
        )
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    if not IMAGE.exists():
        raise FileNotFoundError(IMAGE)
    for current_stem in STEMS:
        update_pptx(current_stem)
        update_html(current_stem)
        export(current_stem)
        print(current_stem)
