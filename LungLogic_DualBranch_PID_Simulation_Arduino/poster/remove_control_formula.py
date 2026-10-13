"""Remove only the formula strip from the 2026-10-02 simulation poster."""

from __future__ import annotations

import base64
import zipfile
from pathlib import Path

from pptx import Presentation

from build_restructured_poster import control_figure


POSTER_DIR = Path(__file__).resolve().parent
SOURCE_STEM = "LungLogic_Poster_2026-10-02_with_simulation"
OUTPUT_STEM = SOURCE_STEM + "_no_formula"
SOURCE_PPTX = POSTER_DIR / f"{SOURCE_STEM}.pptx"
SOURCE_HTML = POSTER_DIR / f"{SOURCE_STEM}.html"
OUTPUT_PPTX = POSTER_DIR / f"{OUTPUT_STEM}.pptx"
OUTPUT_PDF = POSTER_DIR / f"{OUTPUT_STEM}.pdf"
OUTPUT_HTML = POSTER_DIR / f"{OUTPUT_STEM}.html"
OUTPUT_PREVIEW = POSTER_DIR / f"{OUTPUT_STEM}_preview.png"


def find_control_picture():
    deck = Presentation(SOURCE_PPTX)
    if len(deck.slides) != 1:
        raise ValueError("Expected a one-page poster")
    candidates = [
        shape for shape in deck.slides[0].shapes
        if shape.shape_type == 13
        and 28 < shape.top / 914400 < 29
        and shape.left / 914400 > 16
    ]
    if len(candidates) != 1:
        raise ValueError("Could not uniquely identify the control diagram")
    shape = candidates[0]
    rel_id = shape._element.blipFill.blip.rEmbed
    media_name = str(shape.part.related_part(rel_id).partname).lstrip("/")
    return media_name, shape.image.blob


def replace_media(media_name: str, updated_blob: bytes) -> None:
    with zipfile.ZipFile(SOURCE_PPTX) as source, zipfile.ZipFile(
        OUTPUT_PPTX, "w"
    ) as output:
        if media_name not in source.namelist():
            raise ValueError(f"Missing media entry: {media_name}")
        for item in source.infolist():
            data = updated_blob if item.filename == media_name else source.read(item.filename)
            output.writestr(item, data)


def replace_html(original_blob: bytes, updated_blob: bytes) -> None:
    document = SOURCE_HTML.read_text(encoding="utf-8")
    old = base64.b64encode(original_blob).decode("ascii")
    if document.count(old) != 1:
        raise ValueError("Control image did not occur exactly once in the HTML")
    new = base64.b64encode(updated_blob).decode("ascii")
    OUTPUT_HTML.write_text(document.replace(old, new, 1), encoding="utf-8")


def export_pdf_and_preview() -> None:
    import win32com.client

    powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
    powerpoint.Visible = True
    try:
        deck = powerpoint.Presentations.Open(str(OUTPUT_PPTX.resolve()), WithWindow=False)
        deck.SaveAs(str(OUTPUT_PDF.resolve()), 32)
        deck.Slides(1).Export(str(OUTPUT_PREVIEW.resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        powerpoint.Quit()


def main() -> None:
    for path in (SOURCE_PPTX, SOURCE_HTML):
        if not path.exists():
            raise FileNotFoundError(path)
    media_name, original_blob = find_control_picture()
    updated_blob = control_figure(show_formula=False).read_bytes()
    replace_media(media_name, updated_blob)
    replace_html(original_blob, updated_blob)
    export_pdf_and_preview()
    for path in (OUTPUT_PPTX, OUTPUT_PDF, OUTPUT_HTML, OUTPUT_PREVIEW):
        print(path)


if __name__ == "__main__":
    main()
