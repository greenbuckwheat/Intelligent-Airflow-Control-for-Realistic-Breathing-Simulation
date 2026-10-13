"""Build the revised LungLogic A0 poster from the supplied editable PPTX.

The script keeps the visual language of the original poster, replaces its
placeholders with project assets, corrects the control descriptions, creates a
compact single-lung digital-twin validation figure, and exports PPTX/PDF/HTML.
"""

from __future__ import annotations

import base64
import html
import math
import shutil
import sys
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


POSTER_DIR = Path(__file__).resolve().parent
PROJECT_DIR = POSTER_DIR.parent
DOWNLOADS = Path(r"D:\Downloads")
SOURCE_PPTX = DOWNLOADS / "LungLogic_Poster_2026-09-30 (1).pptx"
SCOPE_PPTX = DOWNLOADS / "Robolung Pneumatic System Scoping Presentation.pptx"
OUTPUT_PPTX = POSTER_DIR / "LungLogic_Poster_Revised.pptx"
OUTPUT_PDF = POSTER_DIR / "LungLogic_Poster_Revised.pdf"
OUTPUT_HTML = POSTER_DIR / "LungLogic_Poster_Revised.html"
OUTPUT_PREVIEW = POSTER_DIR / "LungLogic_Poster_Revised_preview.png"
ASSET_DIR = POSTER_DIR / "assets"

NAVY = RGBColor(0x00, 0x06, 0x43)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x00, 0x8C, 0x95)
LIGHT_TEAL = RGBColor(0xE4, 0xF4, 0xF5)
GREY = RGBColor(0x58, 0x68, 0x78)


def set_text(shape, value: str, size: float, *, bold: bool = False,
             italic: bool = False, color: RGBColor = BLACK,
             align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    """Replace text while applying consistent poster typography."""
    frame = shape.text_frame
    frame.clear()
    lines = value.split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        paragraph.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color


def set_objective_card(shape, objective: str, introduction: str,
                       *, objective_size: float = 26,
                       introduction_size: float = 22) -> None:
    """Format an existing text box as an Objective -> introduction card."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_TEAL
    shape.line.color.rgb = TEAL
    shape.line.width = Pt(1.5)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Pt(14)
    frame.margin_right = Pt(14)
    frame.margin_top = Pt(10)
    frame.margin_bottom = Pt(8)

    first = frame.paragraphs[0]
    first.space_after = Pt(6)
    first.line_spacing = 1.0
    label = first.add_run()
    label.text = "OBJECTIVE  |  "
    label.font.name = "Arial"
    label.font.size = Pt(18)
    label.font.bold = True
    label.font.color.rgb = TEAL
    statement = first.add_run()
    statement.text = objective
    statement.font.name = "Arial"
    statement.font.size = Pt(objective_size)
    statement.font.bold = True
    statement.font.color.rgb = NAVY

    second = frame.add_paragraph()
    second.text = introduction
    second.space_before = Pt(0)
    second.space_after = Pt(0)
    second.line_spacing = 1.0
    for run in second.runs:
        run.font.name = "Arial"
        run.font.size = Pt(introduction_size)
        run.font.color.rgb = GREY


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def add_picture_cover(slide, image_path: Path, left, top, width, height):
    """Add an image using centre-crop rather than geometric distortion."""
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    picture = slide.shapes.add_picture(str(image_path), left, top, width, height)
    if image_ratio > box_ratio:
        visible = box_ratio / image_ratio
        picture.crop_left = picture.crop_right = (1.0 - visible) / 2.0
    elif image_ratio < box_ratio:
        visible = image_ratio / box_ratio
        picture.crop_top = picture.crop_bottom = (1.0 - visible) / 2.0
    return picture


def extract_project_assets() -> tuple[Path, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    lung_path = ASSET_DIR / "six_lobe_lung_prototype.png"
    pneumatic_path = ASSET_DIR / "pneumatic_control_case.jpg"
    with zipfile.ZipFile(SCOPE_PPTX) as archive:
        lung_path.write_bytes(archive.read("ppt/media/image12.png"))
        pneumatic_path.write_bytes(archive.read("ppt/media/image13.jpg"))
    return lung_path, pneumatic_path


def build_validation_figure() -> Path:
    sys.path.insert(0, str(PROJECT_DIR))
    from single_lung_pid_gui import (  # pylint: disable=import-outside-toplevel
        STATE_INSPIRATION,
        SingleLungConfig,
        run_single_lung_simulation,
        simulated_inspiratory_tidal_ml,
    )

    cfg = SingleLungConfig()
    cfg.respiratory_rate_bpm = 15.0
    cfg.ie_inspiration = 1.0
    cfg.ie_expiration = 2.0
    cfg.target_volume_ml = 250.0
    cfg.target_fill_time_s = 1.30
    cfg.cycles = 3
    cfg.dead_time_within_cycle = True
    cfg.volume_deadline_guard = True
    cfg.volume_feedback_source = "sensor_integral"
    cfg.pi.kp = 0.010
    cfg.pi.ki = 0.060
    cfg.pi.kd = 0.0
    cfg.pi.integral_limit = 0.5
    cfg.smith.enabled = True
    cfg.smith.delay_s = 0.045
    cfg.smith.correction_gain = 0.6
    result = run_single_lung_simulation(cfg)

    active = result.time_s >= cfg.precondition_s
    time_s = result.time_s[active] - cfg.precondition_s
    tidals = [simulated_inspiratory_tidal_ml(result, cycle) for cycle in range(cfg.cycles)]
    inspiration = result.state == STATE_INSPIRATION
    saturation = float(np.mean(result.opening[inspiration] >= 0.999))

    path = ASSET_DIR / "single_lung_digital_twin_validation.png"
    plt.rcParams.update({"font.family": "Arial", "font.size": 11})
    fig, axes = plt.subplots(1, 3, figsize=(16, 4.0), constrained_layout=True)
    fig.suptitle("Digital-twin pre-validation — nominal single-lung parameters", fontsize=14, fontweight="bold")

    axes[0].plot(time_s, result.volume_ml[active], color="#1f77b4", lw=1.8)
    axes[0].axhline(cfg.target_volume_ml, color="#111111", ls="--", lw=1.1, label="250 mL target")
    axes[0].set(title="Volume", xlabel="Time / s", ylabel="mL", ylim=(-5, 265))
    axes[0].legend(fontsize=9, loc="lower right")

    axes[1].plot(time_s, result.flow_slm[active], color="#167b61", lw=1.4, label="Sensor-model flow")
    axes[1].plot(time_s, result.flow_reference_slm[active], color="#111111", ls="--", lw=1.0, label="Reference")
    axes[1].set(title="Flow tracking", xlabel="Time / s", ylabel="SLM")
    axes[1].legend(fontsize=8, loc="upper right")

    axes[2].plot(time_s, 100.0 * result.opening[active], color="#e07a21", lw=1.5)
    axes[2].set(title="PPV command", xlabel="Time / s", ylabel="Opening / %", ylim=(-3, 105))
    axes[2].text(
        0.03, 0.08,
        f"Mean tidal: {np.mean(tidals):.1f} mL\nFull-open inspiration: {saturation:.0%}",
        transform=axes[2].transAxes,
        fontsize=9,
        bbox={"boxstyle": "round,pad=0.35", "facecolor": "#fff1d6", "edgecolor": "#c99636"},
    )
    for axis in axes:
        axis.grid(True, ls="--", alpha=0.3)
    fig.savefig(path, dpi=220, facecolor="white")
    plt.close(fig)
    return path


def _diagram_card(ax, x: float, y: float, width: float, height: float,
                  title: str, body: str, *, edge: str = "#008c95",
                  face: str = "#e4f4f5") -> None:
    card = FancyBboxPatch(
        (x, y), width, height,
        boxstyle="round,pad=0.012,rounding_size=0.035",
        linewidth=2.2, edgecolor=edge, facecolor=face,
    )
    ax.add_patch(card)
    ax.text(x + width / 2, y + height * 0.66, title, ha="center", va="center",
            fontsize=15, fontweight="bold", color="#17324d")
    ax.text(x + width / 2, y + height * 0.32, body, ha="center", va="center",
            fontsize=10.5, color="#586878", linespacing=1.25)


def build_control_strategy_diagram() -> Path:
    """Create a poster-readable closed-loop diagram instead of dense plant equations."""
    path = ASSET_DIR / "control_strategy_diagram.png"
    plt.rcParams.update({"font.family": "Arial"})
    fig, ax = plt.subplots(figsize=(16, 4.75))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    boxes = [
        (0.015, "Breathing target", "RR · VT · I:E\nfill time"),
        (0.265, "Reference + routing", "Qref(t)\nATM / VAC state machine"),
        (0.515, "Feedforward + PID", "aff from identified model\nPID corrects flow error"),
        (0.765, "Physical path", "PPV · lung · SFM3300\nreal delay + dynamics"),
    ]
    for index, (x, title, body) in enumerate(boxes):
        if index == 2:
            _diagram_card(ax, x, 0.54, 0.215, 0.34, title, body,
                          edge="#f47b20", face="#fff0e2")
        else:
            _diagram_card(ax, x, 0.54, 0.215, 0.34, title, body)
    for x in (0.23, 0.48, 0.73):
        ax.add_patch(FancyArrowPatch((x, 0.71), (x + 0.032, 0.71),
                                     arrowstyle="-|>", mutation_scale=18,
                                     linewidth=2.0, color="#586878"))

    # Measured feedback returns through the optional Smith estimate.
    smith = FancyBboxPatch(
        (0.50, 0.225), 0.24, 0.15,
        boxstyle="round,pad=0.01,rounding_size=0.025",
        linewidth=1.7, edgecolor="#3575c2", facecolor="#e7f0fb",
    )
    ax.add_patch(smith)
    ax.text(0.62, 0.30, "Optional Smith predictor",
            ha="center", va="center", fontsize=12.5, fontweight="bold", color="#17324d")
    ax.text(0.62, 0.255, "identified model + measured delay",
            ha="center", va="center", fontsize=9.8, color="#586878")
    ax.add_patch(FancyArrowPatch((0.875, 0.54), (0.875, 0.30),
                                 arrowstyle="-|>", mutation_scale=17,
                                 linewidth=2.0, color="#008c95"))
    ax.add_patch(FancyArrowPatch((0.875, 0.30), (0.745, 0.30),
                                 arrowstyle="-|>", mutation_scale=17,
                                 linewidth=2.0, color="#008c95"))
    ax.add_patch(FancyArrowPatch((0.50, 0.30), (0.445, 0.30),
                                 arrowstyle="-|>", mutation_scale=17,
                                 linewidth=2.0, color="#008c95"))
    ax.add_patch(FancyArrowPatch((0.445, 0.30), (0.60, 0.54),
                                 connectionstyle="angle3,angleA=180,angleB=90",
                                 arrowstyle="-|>", mutation_scale=17,
                                 linewidth=2.0, color="#008c95"))
    ax.text(0.82, 0.34, "Qmeas(t)", fontsize=10.5, color="#008c95", fontweight="bold")
    ax.text(0.445, 0.34, "Qfeedback", fontsize=10.5, color="#008c95",
            fontweight="bold", ha="center")

    formula = (
        r"$Q_{ref,in}=0.06\,V_T/T_{fill}$"
        r"     $e=Q_{ref}-|Q_{feedback}|$"
        r"     $a=\mathrm{clip}(a_{ff}+K_Pe+K_I\int e\,dt+K_D[-d|Q_{feedback}|/dt],\,0,\,1)$"
    )
    ax.text(0.5, 0.07, formula, ha="center", va="center", fontsize=13.0,
            color="#17324d",
            bbox={"boxstyle": "round,pad=0.5", "facecolor": "#f5f7fa",
                  "edgecolor": "#cbd6df"})
    fig.savefig(path, dpi=220, facecolor="white", bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return path


def build_transfer_workflow_diagram() -> Path:
    """Create the measured-model -> digital-twin -> hardware workflow."""
    path = ASSET_DIR / "model_transfer_workflow.png"
    plt.rcParams.update({"font.family": "Arial"})
    fig, ax = plt.subplots(figsize=(16, 3.8))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    cards = [
        (0.02, "1", "Identify the path", "Full-open ATM/VAC response\nPPV sweep + synchronized delay", "Measured model"),
        (0.35, "2", "Pre-validate digitally", "Test tidal error, tracking,\nsaturation and sensitivity", "Reviewed parameter set"),
        (0.68, "3", "Deploy and lock", "Apply in Developer Mode\nshort conservative hardware run", "Traceable evidence"),
    ]
    palette = [
        ("#008c95", "#e4f4f5"),
        ("#3575c2", "#e7f0fb"),
        ("#f47b20", "#fff0e2"),
    ]
    for (x, number, title, body, outcome), (edge, face) in zip(cards, palette):
        card = FancyBboxPatch(
            (x, 0.12), 0.28, 0.76,
            boxstyle="round,pad=0.016,rounding_size=0.04",
            linewidth=2.2, edgecolor=edge, facecolor=face,
        )
        ax.add_patch(card)
        ax.text(x + 0.038, 0.75, number, ha="center", va="center", fontsize=14,
                fontweight="bold", color="white",
                bbox={"boxstyle": "circle,pad=0.45", "facecolor": edge,
                      "edgecolor": "white", "linewidth": 1.3})
        ax.text(x + 0.075, 0.75, title, ha="left", va="center",
                fontsize=14.5, fontweight="bold", color="#17324d")
        ax.plot([x + 0.03, x + 0.25], [0.61, 0.61], color=edge, lw=1.4)
        ax.text(x + 0.14, 0.46, body, ha="center", va="center",
                fontsize=10.3, color="#43586c", linespacing=1.25)
        ax.text(x + 0.14, 0.20, outcome, ha="center", va="center",
                fontsize=10.8, fontweight="bold", color=edge,
                bbox={"boxstyle": "round,pad=0.32", "facecolor": "white",
                      "edgecolor": "none"})
    for x in (0.305, 0.635):
        ax.add_patch(FancyArrowPatch((x, 0.50), (x + 0.04, 0.50),
                                     arrowstyle="-|>", mutation_scale=18,
                                     linewidth=2.0, color="#586878"))
    fig.savefig(path, dpi=220, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    return path


def revise_pptx(lung_image: Path, pneumatic_image: Path, result_image: Path,
                control_diagram: Path, workflow_diagram: Path) -> None:
    shutil.copy2(SOURCE_PPTX, OUTPUT_PPTX)
    presentation = Presentation(OUTPUT_PPTX)
    slide = presentation.slides[0]
    shapes = list(slide.shapes)

    set_text(shapes[1], "LungLogic: Model-Based Airflow Control", 78, bold=True, color=WHITE)
    set_text(shapes[2], "with Measured-Delay Compensation and Auto PID", 68, bold=True, color=WHITE)
    set_text(
        shapes[3],
        "ME-12-TYIN-007 | Yuqiao Li, Chun Yin Ku (John), Jiaquan Zhong | Supervisors: Jonathan Eden, Ying Tan | Mechatronics Engineering",
        31,
        color=WHITE,
    )

    set_text(shapes[6], "1. Pneumatic System & Requirements", 44, bold=True, color=WHITE)
    set_objective_card(
        shapes[7],
        "Reproduce a physiologic breathing pattern in one 250 mL lung branch.",
        "Design targets combine respiratory timing, tidal volume and repeatability despite valve nonlinearity, sensor dynamics and real pneumatic delay.",
    )
    table = shapes[8].table
    targets = [
        ("Respiratory rate", "12–16 bpm"),
        ("Tidal volume", "250 mL / lung"),
        ("I:E ratio", "1:2"),
        ("Repeatability target", "< ±5%"),
    ]
    for row, (left, right) in enumerate(targets):
        table.cell(row, 0).text = left
        table.cell(row, 1).text = right
        for col in range(2):
            for paragraph in table.cell(row, col).text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(30)
                    run.font.bold = col == 1
    image_box = (shapes[10].left, shapes[10].top, shapes[10].width, shapes[10].height)
    remove_shape(shapes[10])
    remove_shape(shapes[11])
    add_picture_cover(slide, lung_image, *image_box)
    set_text(shapes[12], "Fig. 1. Single-lung pneumatic route and six-lobe breathing phantom.", 28, italic=True)

    set_text(shapes[15], "2. Embedded Hardware Architecture", 44, bold=True, color=WHITE)
    set_objective_card(
        shapes[16],
        "Connect software commands to safe, deterministic pneumatic actuation.",
        "The host configures and logs experiments; the embedded controller executes the 100 Hz state machine and flow loop through a protected 24 V bridge.",
    )
    set_text(shapes[18], "PC / Python GUI", 27, bold=True)
    set_text(shapes[19], "Setpoints, telemetry, Auto PID", 23)
    set_text(shapes[21], "Embedded controller", 27, bold=True)
    set_text(shapes[22], "State machine and flow feedback", 23)
    set_text(shapes[24], "24 V bridge PCB", 27, bold=True)
    set_text(shapes[25], "Power and protected interfaces", 23)
    set_text(shapes[27], "Valves", 27, bold=True)
    set_text(shapes[28], "ATM, vacuum and PPV", 23)
    set_text(shapes[30], "Sensors", 27, bold=True)
    set_text(shapes[31], "Flow feedback and volume estimate", 23)
    set_text(shapes[38], "Configuration / telemetry", 19)
    system_box = (shapes[40].left, shapes[40].top, shapes[40].width, shapes[40].height)
    remove_shape(shapes[40])
    remove_shape(shapes[41])
    add_picture_cover(slide, pneumatic_image, *system_box)
    set_text(shapes[42], "Fig. 2. Host–controller–PCB architecture and pneumatic control case.", 28, italic=True)

    set_text(shapes[45], "3. Control Software & Workflow", 44, bold=True, color=WHITE)
    set_objective_card(
        shapes[46],
        "Provide a visible, adjustable tool for operation, tuning and traceable experiments.",
        "User Mode simplifies routine runs; Developer Mode exposes PID, model, measured-delay and calibration controls.",
        objective_size=24,
        introduction_size=20,
    )
    set_text(shapes[47], "User Mode", 25, bold=True)
    set_text(shapes[48], "Developer Mode", 25, bold=True)
    set_text(shapes[51], "Fig. 3. Telemetry-driven dynamic model and protected developer controls.", 28, italic=True)
    ui_table = shapes[52].table
    functions = [
        ("Breathing setpoints", "Routine experiment setup"),
        ("Dynamic Model", "Live telemetry visualisation"),
        ("PID / Smith / Auto PID", "Controller development"),
        ("Measured parameters", "Locked after identification"),
        ("CSV logging", "Traceable experiment review"),
    ]
    for row, (left, right) in enumerate(functions):
        ui_table.cell(row, 0).text = left
        ui_table.cell(row, 1).text = right
        for col in range(2):
            for paragraph in ui_table.cell(row, col).text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(28)
                    run.font.bold = col == 1
    workflow_box = (shapes[53].left, shapes[53].top, shapes[53].width, shapes[53].height)
    remove_shape(shapes[53])
    add_picture_cover(slide, workflow_diagram, *workflow_box)
    set_text(
        shapes[54],
        "Workflow: identify → simulate → lock/deploy. Measured model and delay values are protected by default; current scope is one 250 mL branch.",
        18,
    )

    set_text(shapes[57], "4. Control Method & Validation", 44, bold=True, color=WHITE)
    set_objective_card(
        shapes[58],
        "Track target flow and tidal volume despite nonlinear PPV response and measured delay.",
        "Feedforward supplies nominal valve opening; measured-flow PID removes residual error; Smith prediction is optional.",
        objective_size=24,
        introduction_size=20,
    )
    control_box = (
        shapes[59].left,
        shapes[59].top,
        shapes[59].width,
        (shapes[62].top + shapes[62].height) - shapes[59].top,
    )
    for shape in shapes[59:63]:
        remove_shape(shape)
    add_picture_cover(slide, control_diagram, *control_box)
    set_text(shapes[63], "Auto PID: bounded run-to-run PI search", 31, bold=True)
    set_text(
        shapes[64],
        "Optional tuning evaluates a stable baseline and bounded Kp/Ki candidates over complete hardware runs. Flow error, tidal error and PPV saturation form the score. Kd and identified model/delay parameters remain fixed; the user reviews the best tested gains.",
        25,
    )
    set_text(shapes[65], "Measured-delay compensation", 31, bold=True)
    set_text(
        shapes[66],
        "Hardware adds no artificial delay. Optional Smith prediction uses the identified first-order model and measured command-to-feedback latency. Re-identify after changing tubing, valves, sensor filtering or sampling.",
        25,
    )
    result_box = (shapes[67].left, shapes[67].top, shapes[67].width, shapes[67].height)
    remove_shape(shapes[67])
    remove_shape(shapes[68])
    add_picture_cover(slide, result_image, *result_box)
    set_text(
        shapes[69],
        "Fig. 4. Digital-twin pre-validation at 15 bpm. Model-based result; measured hardware validation remains required.",
        25,
        italic=True,
    )
    set_text(
        shapes[70],
        "Industry collaborator: 4DMedical | Prototype images: project scope deck (19 March 2026) | Simulation result: current single-lung model | Bench prototype—not for clinical use.",
        15,
    )

    presentation.save(OUTPUT_PPTX)


def rgb_hex(rgb, fallback: str) -> str:
    try:
        return f"#{rgb.rgb}"
    except (AttributeError, TypeError, ValueError):
        return fallback


def export_html() -> None:
    """Create an editable, browser-printable approximation of the PPTX poster."""
    presentation = Presentation(OUTPUT_PPTX)
    slide = presentation.slides[0]
    width_cm = presentation.slide_width / 360000
    height_cm = presentation.slide_height / 360000
    items: list[str] = []

    def pos(shape) -> str:
        return (
            f"left:{shape.left / 360000:.4f}cm;top:{shape.top / 360000:.4f}cm;"
            f"width:{shape.width / 360000:.4f}cm;height:{shape.height / 360000:.4f}cm;"
        )

    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            mime = shape.image.content_type
            data = base64.b64encode(shape.image.blob).decode("ascii")
            items.append(
                f'<img class="item picture" style="{pos(shape)}" '
                f'src="data:{mime};base64,{data}" alt="Poster figure">'
            )
            continue
        if shape.has_table:
            rows = []
            for row in shape.table.rows:
                cells = "".join(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells)
                rows.append(f"<tr>{cells}</tr>")
            items.append(f'<table class="item data-table" style="{pos(shape)}">{"".join(rows)}</table>')
            continue
        if shape.shape_type in (MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.AUTO_SHAPE) and not shape.has_text_frame:
            items.append(f'<div class="item line" style="{pos(shape)}"></div>')
            continue

        fill = "transparent"
        line = "transparent"
        try:
            if shape.fill.type is not None:
                fill = rgb_hex(shape.fill.fore_color, "transparent")
        except (AttributeError, TypeError, ValueError):
            pass
        try:
            line = rgb_hex(shape.line.color, "transparent")
        except (AttributeError, TypeError, ValueError):
            pass
        text_value = html.escape(shape.text) if shape.has_text_frame else ""
        if text_value:
            text_value = text_value.replace("\n", "<br>")
            run = next(
                (run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs),
                None,
            )
            size = run.font.size.pt if run and run.font.size else 24
            bold = "700" if run and run.font.bold else "400"
            italic = "italic" if run and run.font.italic else "normal"
            color = rgb_hex(run.font.color, "#000000") if run else "#000000"
            align = "center" if shape.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER else "left"
            items.append(
                f'<div class="item text" style="{pos(shape)}background:{fill};border:1px solid {line};'
                f'font-size:{size}pt;font-weight:{bold};font-style:{italic};color:{color};text-align:{align};">'
                f'{text_value}</div>'
            )
        else:
            items.append(
                f'<div class="item shape" style="{pos(shape)}background:{fill};border:1px solid {line};"></div>'
            )

    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>LungLogic Revised A0 Poster</title>
<style>
*{{box-sizing:border-box}}html,body{{margin:0;padding:0}}body{{background:#e6e8ec;font-family:Arial,sans-serif}}
.note{{max-width:1000px;margin:14px auto;font-size:14px}}.viewport{{margin:20px auto;position:relative}}
.poster{{width:{width_cm:.4f}cm;height:{height_cm:.4f}cm;position:relative;background:white;transform-origin:top left;overflow:hidden}}
.item{{position:absolute}}.text{{padding:.12cm .18cm;line-height:1.08;overflow:hidden}}.picture{{object-fit:cover;display:block}}
.shape{{padding:0}}.data-table{{border-collapse:collapse;table-layout:fixed;background:white;font-size:27pt}}
.data-table td{{border-bottom:1px solid #c8c8c8;padding:.13cm .18cm;vertical-align:middle}}.data-table td:nth-child(2){{font-weight:700}}
@page{{size:841mm 1189mm;margin:0}}@media print{{html,body{{width:841mm;height:1189mm;background:white}}.note{{display:none}}.viewport{{margin:0!important;width:841mm!important;height:1189mm!important}}.poster{{transform:none!important}}}}
</style></head><body><p class="note">A0 portrait poster. Print at 100% with browser headers and footers disabled.</p>
<main class="viewport"><article class="poster">{''.join(items)}</article></main>
<script>const p=document.querySelector('.poster'),v=document.querySelector('.viewport');function fit(){{const s=Math.min(1,(innerWidth-32)/p.offsetWidth);p.style.transform=`scale(${{s}})`;v.style.width=`${{p.offsetWidth*s}}px`;v.style.height=`${{p.offsetHeight*s}}px`;}}addEventListener('resize',fit);fit();</script>
</body></html>"""
    OUTPUT_HTML.write_text(document, encoding="utf-8")


def export_powerpoint_outputs() -> None:
    import win32com.client  # pylint: disable=import-outside-toplevel

    powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
    powerpoint.Visible = True
    try:
        deck = powerpoint.Presentations.Open(str(OUTPUT_PPTX.resolve()), WithWindow=False)
        deck.SaveAs(str(OUTPUT_PDF.resolve()), 32)
        deck.Slides(1).Export(str(OUTPUT_PREVIEW.resolve()), "PNG", 2384, 3370)
        deck.Close()
    finally:
        powerpoint.Quit()


def main() -> None:
    if not SOURCE_PPTX.exists():
        raise FileNotFoundError(SOURCE_PPTX)
    POSTER_DIR.mkdir(parents=True, exist_ok=True)
    lung_image, pneumatic_image = extract_project_assets()
    result_image = build_validation_figure()
    control_diagram = build_control_strategy_diagram()
    workflow_diagram = build_transfer_workflow_diagram()
    revise_pptx(
        lung_image,
        pneumatic_image,
        result_image,
        control_diagram,
        workflow_diagram,
    )
    export_html()
    export_powerpoint_outputs()
    print(OUTPUT_PPTX)
    print(OUTPUT_PDF)
    print(OUTPUT_HTML)
    print(OUTPUT_PREVIEW)


if __name__ == "__main__":
    main()
