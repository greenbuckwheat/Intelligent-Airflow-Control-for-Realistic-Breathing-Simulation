"""Update the controller label and embedded-control description on both posters."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation


HERE = Path(__file__).resolve().parent
STEMS = (
    "LungLogic_Poster_2026-10-02_prototype_pcb",
    "LungLogic_Poster_2026-10-02_physical_lung_future_work",
)

OLD_LABEL = "Controller"
NEW_LABEL = "MCU"
OLD_COPY = (
    "The controller runs the 100 Hz breathing and flow loop. "
    "The bridge connects logic to the routing valves, PPV and SFM3300 sensor."
)
NEW_COPY = (
    "MCU executes commands from UI and acquires flow data. A custom bridging PCB "
    "integrates 24 V valve drivers, proportional-valve control circuitry, and sensor interfaces."
)


def replace_text_preserving_format(shape, text: str) -> None:
    text_frame = shape.text_frame
    first_run = text_frame.paragraphs[0].runs[0]
    first_run.text = text
    for paragraph in text_frame.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    for run in list(text_frame.paragraphs[0].runs)[1:]:
        run._element.getparent().remove(run._element)


def update_pptx(stem: str) -> None:
    path = HERE / f"{stem}.pptx"
    deck = Presentation(path)
    slide = deck.slides[0]

    labels = [
        shape for shape in slide.shapes
        if shape.has_text_frame and shape.text.strip() in (OLD_LABEL, NEW_LABEL)
    ]
    descriptions = [
        shape for shape in slide.shapes
        if shape.has_text_frame and shape.text.strip() in (OLD_COPY, NEW_COPY)
    ]
    if len(labels) != 1 or len(descriptions) != 1:
        raise ValueError(f"Embedded-control text targets not found in {stem}")

    replace_text_preserving_format(labels[0], NEW_LABEL)
    replace_text_preserving_format(descriptions[0], NEW_COPY)
    # Use the existing gap before the PCB photograph so the longer copy wraps
    # cleanly without changing its font size or overlapping another element.
    descriptions[0].width = round(28.8 * 360000)
    deck.save(path)


def update_html(stem: str) -> None:
    path = HERE / f"{stem}.html"
    document = path.read_text(encoding="utf-8")
    if OLD_LABEL in document:
        document = document.replace(f">{OLD_LABEL}</div>", f">{NEW_LABEL}</div>", 1)
    elif f">{NEW_LABEL}</div>" not in document:
        raise ValueError(f"Controller label not found in HTML: {stem}")

    if OLD_COPY in document:
        document = document.replace(OLD_COPY, NEW_COPY, 1)
    elif NEW_COPY not in document:
        raise ValueError(f"Embedded-control copy not found in HTML: {stem}")

    old_geometry = "left:43.2816cm;top:50.1650cm;width:24.8920cm;height:5.2070cm;"
    new_geometry = "left:43.2816cm;top:50.1650cm;width:28.8000cm;height:5.2070cm;"
    if old_geometry in document:
        document = document.replace(old_geometry, new_geometry, 1)
    elif new_geometry not in document:
        raise ValueError(f"Embedded-control text geometry not found in HTML: {stem}")
    path.write_text(document, encoding="utf-8")


def export(stem: str) -> None:
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{stem}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{stem}.pdf").resolve()), 32)
        deck.Slides(1).Export(
            str((HERE / f"{stem}_preview.png").resolve()), "PNG", 2384, 3370
        )
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    for current_stem in STEMS:
        update_pptx(current_stem)
        update_html(current_stem)
        export(current_stem)
        print(current_stem)
