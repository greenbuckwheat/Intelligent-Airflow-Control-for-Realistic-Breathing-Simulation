"""Expand Results with measured PPV data and remove Future Work.

The first poster row is compacted, Section 3 is shifted upward, and Section 4
is rebuilt as a full-width simulation/measurement results section.
"""

from __future__ import annotations

import base64
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


HERE = Path(__file__).resolve().parent
WORKSPACE = HERE.parents[1]
STEMS = (
    "LungLogic_Poster_2026-10-02_prototype_pcb",
    "LungLogic_Poster_2026-10-02_physical_lung_future_work",
)

SIMULATION = HERE / "assets" / "single_lung_validation_two_panel.png"
MEASURED_FLOW = next(
    WORKSPACE.glob("*/flow_control_exp/analysis/poster_figures/measured_flow_vs_ppv_voltage.png")
)
IDENTIFIED_DYNAMICS = next(
    WORKSPACE.glob("*/flow_control_exp/analysis/poster_figures/identified_time_constant_vs_ppv_voltage.png")
)

NAVY = "000643"
TEXT = "17324D"
BORDER = "C6CEDA"
MUTED = "586878"
FIRST_ROW_HEIGHT = 44.15
SECTION_SHIFT = 4.44
SECTION4_TOP = 87.762
SECTION4_BOTTOM = 117.1956

UI_FIGURE_TOP = 69.15
UI_FIGURE_WIDTH = 18.70
UI_FIGURE_LEFTS = (1.65, 21.65)
UI_CAPTION_TOP = 81.18
AUTO_PID_HEADING_TOP = 83.05
AUTO_PID_COPY_TOP = 84.55

RESULT_SUMMARY = (
    "The measured PPV response qualitatively matches the simulation model: increasing command "
    "voltage attenuates inspiratory flow, and the transient follows a first-order decay."
)

FIGURE_CAPTIONS = (
    ("XV LVAS ventilation map.", "Fig. 1. XV LVAS ventilation map."),
    ("Breathing-cycle illustration at 15 bpm and I:E = 1:2.",
     "Fig. 2. Breathing-cycle illustration at 15 bpm and I:E = 1:2."),
    ("Fig. 2. Single-lung pneumatic route and 250 mL prototype.",
     "Fig. 3. Single-lung pneumatic route and 250 mL prototype."),
    ("Fig. 3. Control architecture and bridge PCB connected to STM32.",
     "Fig. 4. Control architecture and bridge PCB connected to STM32."),
    ("Fig. 4. User and Developer modes.", "Fig. 5. User and Developer modes."),
    ("Fig. 5. Feedback control with optional Smith prediction.",
     "Fig. 6. Feedback control with optional Smith prediction."),
)


def cm(value: float) -> int:
    return round(value * 360000)


def color(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def replace_text_preserving_format(target, text: str) -> None:
    text_frame = target.text_frame
    first_run = text_frame.paragraphs[0].runs[0]
    first_run.text = text
    for paragraph in text_frame.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    for run in list(text_frame.paragraphs[0].runs)[1:]:
        run._element.getparent().remove(run._element)


def add_box(slide, left, top, width, height, fill, line) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, cm(left), cm(top), cm(width), cm(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line)
    shape.line.width = Pt(1)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size,
    *,
    bold=False,
    italic=False,
    font_color=TEXT,
    align=PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(cm(left), cm(top), cm(width), cm(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = cm(0.18)
    frame.margin_right = cm(0.18)
    frame.margin_top = cm(0.12)
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = 0
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(font_color)


def find_text(slide, text: str):
    return [
        shape for shape in slide.shapes
        if shape.has_text_frame and shape.text.strip() == text
    ]


def update_pptx(stem: str) -> None:
    path = HERE / f"{stem}.pptx"
    deck = Presentation(path)
    slide = deck.slides[0]

    header_old = (
        "ME-12-TYIN-007 | Yuqiao Li, Chun Yin Ku (John), Jiaquan Zhong | "
        "Supervisors: Jonathan Eden, Ying Tan | Mechatronics Engineering"
    )
    header_new = (
        "ME-12-TYIN-007 | Yuqiao Li, Chun Yin Ku (John), Jiaquan Zhong | "
        "Supervisors: Jonathan Eden, Ying Tan"
    )
    headers = find_text(slide, header_old) or find_text(slide, header_new)
    if len(headers) != 1:
        raise ValueError(f"Poster header not found in {stem}")
    replace_text_preserving_format(headers[0], header_new)

    # Compact the aligned Introduction and Hardware panels.
    for shape in slide.shapes:
        left = shape.left / 360000
        top = shape.top / 360000
        if abs(top - 13.0048) < 0.01 and (
            abs(left - 0.7112) < 0.01 or abs(left - 42.1894) < 0.01
        ) and shape.height / 360000 > 40:
            shape.height = cm(FIRST_ROW_HEIGHT)

    tables = [shape for shape in slide.shapes if shape.has_table]
    if len(tables) != 1:
        raise ValueError(f"Expected one Introduction table in {stem}")
    table_shape = tables[0]
    table_shape.height = cm(6.2)
    for row in table_shape.table.rows:
        row.height = cm(1.55)
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(27)
    table_shape.table.cell(1, 0).text_frame.paragraphs[0].runs[0].text = "Inhale:Exhale time"

    note = find_text(slide, "Nominal reference values shown; operating parameters are programmable.")
    source = find_text(slide, "Source: 4DMedical, XV LVAS Info Sheet, DOC-5070 US V1 (2023), pp. 1-2.")
    if len(note) != 1 or len(source) != 1:
        raise ValueError(f"Introduction note/source not found in {stem}")
    note[0].top, note[0].height = cm(53.85), cm(1.2)
    source[0].top, source[0].height = cm(55.25), cm(1.2)

    objective_old = (
        "Commercial lung models cannot meet 4DMedical's needs. Our programmable airflow system "
        "will drive a patient-realistic physical lung developed by the medical team. This stage "
        "validates one 250 mL branch."
    )
    objective_previous = (
        "Commercial lung models cannot accommodate specific research and development requirements. "
        "Our programmable airflow system will drive a patient-realistic physical lung developed by "
        "the medical team. This stage validates one 250 mL branch."
    )
    objective_with_scanner_old = (
        "Commercial lung models cannot accommodate specific R&D requirements. Together with "
        "4DMedical's XV scanner, our programmable airflow system and the medical team's "
        "patient-realistic lung model enable controlled ventilation-imaging research. This stage "
        "validates one 250 mL branch."
    )
    objective_with_scanner_previous = (
        "Commercial lung models cannot meet specific R&D requirements. With 4DMedical's XV scanner, "
        "our airflow system and the medical team's patient-realistic lung model enable controlled "
        "ventilation-imaging research. This stage validates one 250 mL branch."
    )
    objective_abbreviated = (
        "Commercial lung models cannot meet specific R&D requirements. Paired with 4DMedical's XV "
        "scanner, our airflow controller and the medical team's patient-realistic lung model support "
        "controlled ventilation-imaging studies."
    )
    objective_full_previous = (
        "Commercial lung models cannot meet specific research and development needs. Using "
        "4DMedical's XV scanner, our airflow controller and the medical team's lung model support "
        "controlled ventilation-imaging studies. This stage validates one 250 mL branch prototype."
    )
    objective_full_previous_2 = (
        "Commercial lung models cannot meet specific research and development needs. With its XV "
        "scanner, 4DMedical can use our airflow controller and the medical team's lung model in "
        "ventilation-imaging studies. This stage validates one 250 mL branch prototype."
    )
    objective_full_previous_3 = (
        "Commercial models cannot meet specific research and development needs. With its XV scanner, "
        "4DMedical can use our controller and the medical team's lung model in "
        "ventilation-imaging studies. This stage validates one 250 mL branch prototype."
    )
    objective_full_previous_4 = (
        "Commercial models cannot meet specific research and development needs. With its XV scanner, "
        "4DMedical can use our control system and the medical team's lung model for ventilation "
        "research. This stage validates one 250 mL branch prototype."
    )
    objective_full_previous_5 = (
        "Commercial models cannot meet specific research and development needs. 4DMedical can use "
        "its XV scanner with our controller and the medical team's lung model for ventilation research. "
        "This stage validates one 250 mL branch prototype."
    )
    objective_new = (
        "Commercial models cannot meet specific research and development needs. 4DMedical's XV "
        "scanner enables research using our controller and the medical team's lung model. "
        "This stage validates one 250 mL branch prototype."
    )
    objectives = (
        find_text(slide, objective_old)
        or find_text(slide, objective_previous)
        or find_text(slide, objective_with_scanner_old)
        or find_text(slide, objective_with_scanner_previous)
        or find_text(slide, objective_abbreviated)
        or find_text(slide, objective_full_previous)
        or find_text(slide, objective_full_previous_2)
        or find_text(slide, objective_full_previous_3)
        or find_text(slide, objective_full_previous_4)
        or find_text(slide, objective_full_previous_5)
        or find_text(slide, objective_new)
    )
    if len(objectives) != 1:
        raise ValueError(f"Introduction objective not found in {stem}")
    replace_text_preserving_format(objectives[0], objective_new)

    # Delete the Hardware objective and renumber the remaining subsections.
    for old in (
        "2.1 Objective",
        "Control one lung branch and measure airflow for closed-loop regulation and calibration.",
    ):
        for shape in find_text(slide, old):
            remove_shape(shape)
    renames = {
        "2.2 Pneumatic system prototype": "2.1 Pneumatic system prototype",
        "2.3 Embedded control architecture": "2.2 Embedded control architecture",
    }
    for old, new in renames.items():
        matches = find_text(slide, old) or find_text(slide, new)
        if len(matches) != 1:
            raise ValueError(f"Hardware subsection not found in {stem}: {old}")
        replace_text_preserving_format(matches[0], new)

    hardware_top_map = {
        21.7424: 15.9512,
        23.8252: 18.0000,
        23.4950: 17.6700,
        42.5704: 36.7500,
        44.8564: 39.0500,
        46.7360: 40.9860,
        47.3710: 41.6210,
        47.4218: 41.6718,
        50.1650: 44.4150,
        50.3936: 44.6436,
        58.8010: 54.8500,
    }
    for shape in slide.shapes:
        if shape.left / 360000 < 42:
            continue
        current_top = shape.top / 360000
        for old_top, new_top in hardware_top_map.items():
            if abs(current_top - old_top) < 0.01:
                shape.top = cm(new_top)
                break

    # Reclaim the saved height by moving Section 3 upward once.
    section3_titles = find_text(slide, "3. Integrated Control System & Software")
    if len(section3_titles) != 1:
        raise ValueError(f"Section 3 title not found in {stem}")
    if section3_titles[0].top / 360000 > 61:
        for shape in slide.shapes:
            top = shape.top / 360000
            if 62.8 <= top < 92.1:
                shape.top = round(shape.top - cm(SECTION_SHIFT))
        old_results_start = 92.1
    else:
        old_results_start = 87.6

    # Enlarge the two UI screenshots and use the previously empty space below
    # them.  Keep the right-hand feedback diagram unchanged.
    ui_headings = {
        "User Mode": (UI_FIGURE_LEFTS[0], 18.70),
        "Developer Mode": (UI_FIGURE_LEFTS[1], 18.70),
    }
    for label, (left, width) in ui_headings.items():
        matches = find_text(slide, label)
        if len(matches) != 1:
            raise ValueError(f"UI heading not found in {stem}: {label}")
        matches[0].left = cm(left)
        matches[0].width = cm(width)

    ui_pictures = sorted(
        (
            shape for shape in slide.shapes
            if shape.shape_type == 13
            and shape.left / 360000 < 42
            and 68.5 < shape.top / 360000 < 70.0
        ),
        key=lambda shape: shape.left,
    )
    if len(ui_pictures) != 2:
        raise ValueError(f"Expected two UI screenshots in {stem}")
    for picture, left in zip(ui_pictures, UI_FIGURE_LEFTS):
        aspect = picture.image.size[0] / picture.image.size[1]
        picture.left = cm(left)
        picture.top = cm(UI_FIGURE_TOP)
        picture.width = cm(UI_FIGURE_WIDTH)
        picture.height = cm(UI_FIGURE_WIDTH / aspect)

    fig3 = (
        find_text(slide, "Fig. 3. User and Developer modes.")
        or find_text(slide, "Fig. 4. User and Developer modes.")
        or find_text(slide, "Fig. 5. User and Developer modes.")
    )
    auto_pid_heading = find_text(slide, "3.4 Optional Auto PID and measured-delay compensation")
    if len(fig3) != 1 or len(auto_pid_heading) != 1:
        raise ValueError(f"Section 3 labels not found in {stem}")
    fig3[0].top = cm(UI_CAPTION_TOP)
    auto_pid_heading[0].top = cm(AUTO_PID_HEADING_TOP)
    auto_pid_copy = [
        shape for shape in slide.shapes
        if shape.has_text_frame
        and shape.text.startswith("Auto PID scores bounded Kp/Ki")
    ]
    if len(auto_pid_copy) != 1:
        raise ValueError(f"Auto PID copy not found in {stem}")
    auto_pid_copy[0].top = cm(AUTO_PID_COPY_TOP)

    # Number every visual group consecutively and apply one caption style.
    for old_caption, new_caption in FIGURE_CAPTIONS:
        matches = find_text(slide, old_caption) or find_text(slide, new_caption)
        if len(matches) != 1:
            raise ValueError(f"Figure caption not found in {stem}: {old_caption}")
        replace_text_preserving_format(matches[0], new_caption)
        for paragraph in matches[0].text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(28)
                run.font.italic = True

    # Remove the complete old Results/Future Work section and rebuild it.
    for shape in list(slide.shapes):
        top = shape.top / 360000
        if old_results_start <= top < 117.3:
            remove_shape(shape)

    add_box(
        slide, 0.7112, SECTION4_TOP, 82.6770,
        SECTION4_BOTTOM - SECTION4_TOP, "FFFFFF", BORDER,
    )
    add_box(slide, 0.7112, SECTION4_TOP, 82.6770, 2.3876, NAVY, NAVY)
    add_text(
        slide, 1.1684, SECTION4_TOP + 0.3048, 81.7626, 1.7780,
        "4. Results", 48, bold=True, font_color="FFFFFF",
    )
    add_text(slide, 1.30, 90.38, 35.0, 1.63, "4.1 Digital-twin simulation", 36, bold=True)
    add_text(slide, 36.70, 90.38, 45.7, 1.63, "4.2 Physical measurements", 36, bold=True)

    slide.shapes.add_picture(
        str(SIMULATION), cm(1.30), cm(93.20), width=cm(35.0), height=cm(11.105)
    )
    slide.shapes.add_picture(
        str(MEASURED_FLOW), cm(36.80), cm(92.15), width=cm(22.6), height=cm(13.035)
    )
    slide.shapes.add_picture(
        str(IDENTIFIED_DYNAMICS), cm(59.90), cm(92.05), width=cm(22.6), height=cm(13.298)
    )
    add_text(
        slide, 1.30, 105.55, 35.0, 1.70,
        "Fig. 7. Digital-twin volume and flow tracking at 15 bpm.",
        28, italic=True,
    )
    add_text(
        slide, 36.70, 105.55, 45.7, 1.70,
        "Fig. 8. Measured PPV flow response and identified first-order dynamics.",
        28, italic=True,
    )
    add_text(slide, 1.30, 107.55, 81.1, 3.60, RESULT_SUMMARY, 36)

    footer_old = (
        "Industry collaborator: 4DMedical | Background: XV LVAS Info Sheet, DOC-5070 US V1 (2023) | "
        "Control: current single-lung system | Experimental results: pending | Bench research prototype."
    )
    footer_new = (
        "Industry collaborator: 4DMedical | Background: XV LVAS Info Sheet, DOC-5070 US V1 (2023) | "
        "Control: current single-lung system | Results: simulation + benchtop measurements | Bench research prototype."
    )
    footers = find_text(slide, footer_old) or find_text(slide, footer_new)
    if len(footers) != 1:
        raise ValueError(f"Footer not found in {stem}")
    replace_text_preserving_format(footers[0], footer_new)
    deck.save(path)


def style_number(style: str, name: str) -> float | None:
    match = re.search(rf"{name}:([0-9.]+)cm", style)
    return float(match.group(1)) if match else None


def set_style_value(style: str, name: str, value: float) -> str:
    return re.sub(rf"{name}:[0-9.]+cm", f"{name}:{value:.4f}cm", style, count=1)


def transform_styled_tags(document: str, transform) -> str:
    pattern = re.compile(
        r'<div\b[^>]*style="[^"]*"[^>]*>.*?</div>'
        r'|<img\b[^>]*style="[^"]*"[^>]*/>'
        r'|<table\b[^>]*style="[^"]*"[^>]*>.*?</table>',
        flags=re.DOTALL,
    )

    def replace(match):
        tag = match.group(0)
        style_match = re.search(r'style="([^"]*)"', tag)
        if not style_match:
            return tag
        new_tag = transform(tag, style_match.group(1))
        return new_tag

    return pattern.sub(replace, document)


def remove_div_with_text(document: str, text: str) -> str:
    pattern = re.compile(
        r'<div class="item text" style="[^"]*">' + re.escape(text) + r'</div>'
    )
    return pattern.sub("", document, count=1)


def image_html(path: Path, alt: str, left, top, width, height) -> str:
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return (
        f'<img alt="{alt}" class="item picture" '
        f'src="data:image/png;base64,{encoded}" '
        f'style="left:{left:.4f}cm;top:{top:.4f}cm;'
        f'width:{width:.4f}cm;height:{height:.4f}cm;"/>'
    )


def text_html(left, top, width, height, text, size, *, bold=False, italic=False, text_color=TEXT) -> str:
    return (
        f'<div class="item text" style="left:{left:.4f}cm;top:{top:.4f}cm;'
        f'width:{width:.4f}cm;height:{height:.4f}cm;background:transparent;'
        f'border:1px solid transparent;font-size:{size:.1f}pt;'
        f'font-weight:{700 if bold else 400};font-style:{"italic" if italic else "normal"};'
        f'color:#{text_color};text-align:left;">{text}</div>'
    )


def update_html(stem: str) -> None:
    path = HERE / f"{stem}.html"
    document = path.read_text(encoding="utf-8")
    document = document.replace(" | Mechatronics Engineering", "", 1)
    document = remove_div_with_text(document, "2.1 Objective")
    document = remove_div_with_text(
        document,
        "Control one lung branch and measure airflow for closed-loop regulation and calibration.",
    )
    document = document.replace("2.2 Pneumatic system prototype", "2.1 Pneumatic system prototype", 1)
    document = document.replace("2.3 Embedded control architecture", "2.2 Embedded control architecture", 1)
    document = document.replace("Inhale:exhale time", "Inhale:Exhale time", 1)
    document = document.replace(
        "Commercial models cannot meet specific research and development needs. 4DMedical can use "
        "its XV scanner with our controller and the medical team's lung model for ventilation research. "
        "This stage validates one 250 mL branch prototype.",
        "Commercial models cannot meet specific research and development needs. 4DMedical's XV "
        "scanner enables research using our controller and the medical team's lung model. "
        "This stage validates one 250 mL branch prototype.",
        1,
    )
    for old_caption, new_caption in FIGURE_CAPTIONS:
        document = document.replace(old_caption, new_caption, 1)

    hardware_top_map = {
        21.7424: 15.9512, 23.8252: 18.0000, 23.4950: 17.6700,
        42.5704: 36.7500, 44.8564: 39.0500, 46.7360: 40.9860,
        47.3710: 41.6210, 47.4218: 41.6718, 50.1650: 44.4150,
        50.3936: 44.6436, 58.8010: 54.8500,
    }

    # Determine whether Section 3 still needs to move before transforming tags.
    section3_match = re.search(
        r'<div class="item text" style="[^"]*top:([0-9.]+)cm;[^"]*">'
        r'3\. Integrated Control System &amp; Software</div>',
        document,
    )
    if not section3_match:
        raise ValueError(f"Section 3 title not found in HTML: {stem}")
    shift_section3 = float(section3_match.group(1)) > 61
    old_results_start = 92.1 if shift_section3 else 87.6

    def transform(tag: str, style: str) -> str:
        left = style_number(style, "left")
        top = style_number(style, "top")
        height = style_number(style, "height")
        if left is None or top is None:
            return tag
        # Remove the complete old/new Section 4; it is rebuilt below.
        if old_results_start <= top < 117.3:
            return ""
        new_style = style
        if abs(top - 13.0048) < 0.01 and height and height > 40 and (
            abs(left - 0.7112) < 0.01 or abs(left - 42.1894) < 0.01
        ):
            new_style = set_style_value(new_style, "height", FIRST_ROW_HEIGHT)
        if left > 42:
            for old_top, new_top in hardware_top_map.items():
                if abs(top - old_top) < 0.01:
                    new_style = set_style_value(new_style, "top", new_top)
                    break
        if shift_section3 and 62.8 <= top < 92.1:
            new_style = set_style_value(new_style, "top", top - SECTION_SHIFT)
        if tag.startswith('<table') and abs(left - 1.5494) < 0.01:
            new_style = set_style_value(new_style, "height", 6.2)
        if "Nominal reference values shown; operating parameters are programmable." in tag:
            new_style = set_style_value(new_style, "top", 53.85)
            new_style = set_style_value(new_style, "height", 1.2)
        if "Source: 4DMedical, XV LVAS Info Sheet" in tag:
            new_style = set_style_value(new_style, "top", 55.25)
            new_style = set_style_value(new_style, "height", 1.2)
        if ">User Mode</div>" in tag:
            new_style = set_style_value(new_style, "left", UI_FIGURE_LEFTS[0])
            new_style = set_style_value(new_style, "width", UI_FIGURE_WIDTH)
        if ">Developer Mode</div>" in tag:
            new_style = set_style_value(new_style, "left", UI_FIGURE_LEFTS[1])
            new_style = set_style_value(new_style, "width", UI_FIGURE_WIDTH)
        if tag.startswith("<img") and left < 42 and 68.5 < top < 70.0:
            image_left = UI_FIGURE_LEFTS[0] if left < 21 else UI_FIGURE_LEFTS[1]
            image_aspect = 1346 / 856 if image_left == UI_FIGURE_LEFTS[0] else 1354 / 864
            new_style = set_style_value(new_style, "left", image_left)
            new_style = set_style_value(new_style, "top", UI_FIGURE_TOP)
            new_style = set_style_value(new_style, "width", UI_FIGURE_WIDTH)
            new_style = set_style_value(new_style, "height", UI_FIGURE_WIDTH / image_aspect)
        if ("Fig. 3. User and Developer modes." in tag
                or "Fig. 4. User and Developer modes." in tag
                or "Fig. 5. User and Developer modes." in tag):
            new_style = set_style_value(new_style, "top", UI_CAPTION_TOP)
        if "3.4 Optional Auto PID and measured-delay compensation" in tag:
            new_style = set_style_value(new_style, "top", AUTO_PID_HEADING_TOP)
        if "Auto PID scores bounded Kp/Ki" in tag:
            new_style = set_style_value(new_style, "top", AUTO_PID_COPY_TOP)
        if any(new_caption in tag for _, new_caption in FIGURE_CAPTIONS):
            new_style = re.sub(r"font-size:[0-9.]+pt", "font-size:28.0pt", new_style, count=1)
            new_style = re.sub(r"font-style:[^;]+", "font-style:italic", new_style, count=1)
        return tag.replace(f'style="{style}"', f'style="{new_style}"', 1)

    document = transform_styled_tags(document, transform)

    footer_old = (
        "Industry collaborator: 4DMedical | Background: XV LVAS Info Sheet, DOC-5070 US V1 (2023) | "
        "Control: current single-lung system | Experimental results: pending | Bench research prototype."
    )
    footer_new = (
        "Industry collaborator: 4DMedical | Background: XV LVAS Info Sheet, DOC-5070 US V1 (2023) | "
        "Control: current single-lung system | Results: simulation + benchtop measurements | Bench research prototype."
    )
    document = document.replace(footer_old, footer_new, 1)

    result_html = "".join((
        f'<div class="item shape" style="left:0.7112cm;top:{SECTION4_TOP:.4f}cm;'
        f'width:82.6770cm;height:{SECTION4_BOTTOM - SECTION4_TOP:.4f}cm;'
        f'background:#FFFFFF;border:1px solid #{BORDER};"></div>',
        f'<div class="item shape" style="left:0.7112cm;top:{SECTION4_TOP:.4f}cm;'
        f'width:82.6770cm;height:2.3876cm;background:#{NAVY};border:1px solid #{NAVY};"></div>',
        text_html(1.1684, SECTION4_TOP + 0.3048, 81.7626, 1.7780, "4. Results", 48, bold=True, text_color="FFFFFF"),
        text_html(1.30, 90.38, 35.0, 1.63, "4.1 Digital-twin simulation", 36, bold=True),
        text_html(36.70, 90.38, 45.7, 1.63, "4.2 Physical measurements", 36, bold=True),
        image_html(SIMULATION, "Digital-twin volume and flow tracking", 1.30, 93.20, 35.0, 11.105),
        image_html(MEASURED_FLOW, "Measured inspiration flow at different PPV voltages", 36.80, 92.15, 22.6, 13.035),
        image_html(IDENTIFIED_DYNAMICS, "Identified first-order physical-branch dynamics", 59.90, 92.05, 22.6, 13.298),
        text_html(1.30, 105.55, 35.0, 1.70, "Fig. 7. Digital-twin volume and flow tracking at 15 bpm.", 28, italic=True),
        text_html(36.70, 105.55, 45.7, 1.70, "Fig. 8. Measured PPV flow response and identified first-order dynamics.", 28, italic=True),
        text_html(1.30, 107.55, 81.1, 3.60, RESULT_SUMMARY, 36),
    ))
    footer_marker = re.search(
        r'(?=<div class="item text" style="left:0\.8890cm;top:117\.6782cm;)',
        document,
    )
    if not footer_marker:
        raise ValueError(f"Footer marker not found in HTML: {stem}")
    insert_at = footer_marker.start()
    document = document[:insert_at] + result_html + document[insert_at:]
    path.write_text(document, encoding="utf-8")


def export(stem: str) -> None:
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{stem}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{stem}.pdf").resolve()), 32)
        deck.Slides(1).Export(
            str((HERE / f"{stem}_preview.png").resolve()), "PNG", 2384, 3370
        )
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    for required in (SIMULATION, MEASURED_FLOW, IDENTIFIED_DYNAMICS):
        if not required.exists():
            raise FileNotFoundError(required)
    for current_stem in STEMS:
        update_pptx(current_stem)
        update_html(current_stem)
        export(current_stem)
        print(current_stem)
