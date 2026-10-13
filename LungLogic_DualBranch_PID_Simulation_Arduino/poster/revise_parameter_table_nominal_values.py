"""Use full parameter names and restore nominal values in the poster table."""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation


HERE = Path(__file__).resolve().parent
STEMS = (
    "LungLogic_Poster_2026-10-02_prototype_pcb",
    "LungLogic_Poster_2026-10-02_physical_lung_future_work",
)

ROWS = (
    ("Respiratory rate", "12–16 bpm; nominal 15 bpm"),
    ("Inhale:exhale time", "1:2; 1.33 s / 2.67 s at 15 bpm"),
    ("Tidal volume", "250 mL per branch"),
    ("Regional airflow", "Independently programmable by branch"),
)

OLD_NOTES = (
    "Parameters are adjustable to reproduce controlled patient-like breathing conditions.",
    "Nominal design values, not measured performance. Target: 250 mL per branch and VT error within ±5%.",
)
NEW_NOTE = "Nominal reference values shown; operating parameters are programmable."


def replace_text_preserving_format(target, text: str) -> None:
    text_frame = target.text_frame
    first_run = text_frame.paragraphs[0].runs[0]
    first_run.text = text
    for paragraph in text_frame.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    for run in list(text_frame.paragraphs[0].runs)[1:]:
        run._element.getparent().remove(run._element)


def update_pptx(stem: str) -> None:
    path = HERE / f"{stem}.pptx"
    deck = Presentation(path)
    slide = deck.slides[0]
    table_shapes = [shape for shape in slide.shapes if shape.has_table]
    if len(table_shapes) != 1:
        raise ValueError(f"Expected one parameter table in {stem}")
    table = table_shapes[0].table
    if len(table.rows) != 4 or len(table.columns) != 2:
        raise ValueError(f"Unexpected parameter table dimensions in {stem}")
    for row, values in zip(table.rows, ROWS):
        for cell, value in zip(row.cells, values):
            replace_text_preserving_format(cell, value)

    notes = [
        shape for shape in slide.shapes
        if shape.has_text_frame
        and (shape.text.strip() in OLD_NOTES or shape.text.strip() == NEW_NOTE)
    ]
    if len(notes) != 1:
        raise ValueError(f"Parameter note not found in {stem}")
    replace_text_preserving_format(notes[0], NEW_NOTE)
    deck.save(path)


def update_html(stem: str) -> None:
    path = HERE / f"{stem}.html"
    document = path.read_text(encoding="utf-8")
    table_pattern = re.compile(
        r'<table class="item data-table" '
        r'style="left:1\.5494cm;top:47\.4000cm;width:39\.0652cm;height:9\.6520cm;">'
        r'.*?</table>',
        flags=re.DOTALL,
    )
    rows = "".join(
        f"<tr><td>{label}</td><td>{value}</td></tr>" for label, value in ROWS
    )
    replacement = (
        '<table class="item data-table" '
        'style="left:1.5494cm;top:47.4000cm;width:39.0652cm;height:9.6520cm;">'
        + rows
        + "</table>"
    )
    document, count = table_pattern.subn(lambda _: replacement, document, count=1)
    if count != 1:
        raise ValueError(f"Parameter table not found in HTML: {stem}")

    if NEW_NOTE not in document:
        for old_note in OLD_NOTES:
            if old_note in document:
                document = document.replace(old_note, NEW_NOTE, 1)
                break
        else:
            raise ValueError(f"Parameter note not found in HTML: {stem}")
    path.write_text(document, encoding="utf-8")


def export(stem: str) -> None:
    import win32com.client

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = True
    try:
        deck = app.Presentations.Open(str((HERE / f"{stem}.pptx").resolve()), WithWindow=False)
        deck.SaveAs(str((HERE / f"{stem}.pdf").resolve()), 32)
        deck.Slides(1).Export(
            str((HERE / f"{stem}_preview.png").resolve()), "PNG", 2384, 3370
        )
        deck.Close()
    finally:
        app.Quit()


if __name__ == "__main__":
    for current_stem in STEMS:
        update_pptx(current_stem)
        update_html(current_stem)
        export(current_stem)
        print(current_stem)
