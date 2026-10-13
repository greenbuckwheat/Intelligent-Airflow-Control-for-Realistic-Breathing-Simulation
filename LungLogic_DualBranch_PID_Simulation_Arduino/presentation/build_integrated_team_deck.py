"""Extend the team deck with an agenda, three section pages and a new Control UI section."""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


SOURCE = Path(r"D:\Downloads\23-09-Presentation_new.pptx")
HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "23-09-Presentation_Integrated_Control_UI.pptx"

NAVY = RGBColor(18, 45, 70)
NAVY_2 = RGBColor(27, 67, 91)
TEAL = RGBColor(15, 166, 169)
TEAL_DARK = RGBColor(0, 132, 137)
TEAL_LIGHT = RGBColor(229, 246, 246)
BLUE = RGBColor(55, 112, 181)
BLUE_LIGHT = RGBColor(232, 241, 250)
ORANGE = RGBColor(234, 126, 41)
ORANGE_LIGHT = RGBColor(255, 241, 227)
GREEN = RGBColor(33, 145, 101)
RED = RGBColor(202, 65, 55)
INK = RGBColor(29, 43, 57)
MID = RGBColor(91, 106, 119)
LIGHT = RGBColor(246, 248, 250)
LINE = RGBColor(215, 224, 231)
WHITE = RGBColor(255, 255, 255)


def text(slide, value, x, y, w, h, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    frame.vertical_anchor = valign
    for i, line_value in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line_value
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
        p.space_after = Pt(3)
    return shape


def box(slide, x, y, w, h, fill=LIGHT, border=None, rounded=True):
    kind = (MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded
            else MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(1.1)
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def circle(slide, value, x, y, diameter=.38, fill=TEAL_DARK, size=13):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = WHITE
    shape.line.width = Pt(1.5)
    text(slide, str(value), x, y + .01, diameter, diameter - .02, size, True,
         WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return shape


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_header(slide, title):
    text(slide, title, .36, .23, 11.7, .50, 25, True, NAVY)
    box(slide, .36, .80, 1.18, .04, TEAL, None, False)
    text(slide, "4DMedical / LungLogic", 11.55, .30, 1.35, .18, 7.5,
         False, TEAL_DARK, PP_ALIGN.RIGHT)


def takeaway(slide, value):
    box(slide, .45, 6.91, 12.42, .34, TEAL_LIGHT, None, False)
    text(slide, value, .63, 6.98, 12.0, .19, 9.5, True, NAVY,
         valign=MSO_ANCHOR.MIDDLE)


def transition(prs, number, title_value, subtitle):
    slide = blank(prs)
    box(slide, 0, 0, 13.333, 7.5, NAVY, None, False)
    box(slide, 0, 0, .13, 7.5, TEAL, None, False)
    text(slide, number, .70, 1.25, 1.15, .70, 31, True, TEAL)
    text(slide, title_value, 2.12, 1.26, 9.8, .72, 30, True, WHITE)
    line(slide, 2.12, 2.36, 11.95, 2.36, RGBColor(78, 109, 132), 1.3)
    text(slide, subtitle, 2.12, 2.80, 9.6, 1.20, 20, False,
         RGBColor(205, 220, 231), valign=MSO_ANCHOR.MIDDLE)
    return slide


def agenda_slide(prs):
    slide = blank(prs)
    content_header(slide, "Presentation Overview")
    sections = [
        ("01", "Mechanical Part", "Pneumatic architecture\nAirflow sizing\nPhysical prototype", TEAL_LIGHT, TEAL_DARK),
        ("02", "Electrical Part", "Power and bridge board\nSensor interfaces\nSignal verification", BLUE_LIGHT, BLUE),
        ("03", "Control Software & UI", "Python interface\nReusable model + Auto PID\nHardware migration", ORANGE_LIGHT, ORANGE),
    ]
    for i, (num, title_value, body, fill, color) in enumerate(sections):
        x = .70 + i * 4.15
        box(slide, x, 1.52, 3.70, 4.60, fill, color)
        circle(slide, num, x + .28, 1.84, .58, color, 16)
        text(slide, title_value, x + .28, 2.67, 3.10, .47, 20, True, NAVY)
        line(slide, x + .28, 3.30, x + 3.40, 3.30, color, 1.4)
        text(slide, body, x + .30, 3.68, 3.05, 1.45, 15.5, False, NAVY)
    takeaway(slide, "Three parts, one system: pneumatic hardware, electrical interface, and reusable control software.")
    notes(slide, "The presentation has three parts. We first cover the mechanical system, then the electrical interface and signal verification, and finally the Python control software and user interface.")
    return slide


def objective_slide(prs):
    slide = blank(prs)
    content_header(slide, "Objective & requirements: a visible, adjustable Python control tool")
    box(slide, .55, 1.20, 12.20, 1.15, TEAL_LIGHT, TEAL_DARK)
    text(slide, "OBJECTIVE", .83, 1.48, 1.35, .25, 10.5, True, TEAL_DARK)
    text(slide, "Visualise each breath and let users adjust the operating target without changing control code.",
         2.22, 1.36, 9.95, .62, 19, True, NAVY, valign=MSO_ANCHOR.MIDDLE)

    requirements = [
        ("Breathing", "Rate · tidal volume\nI:E ratio · flow timing", TEAL_LIGHT, TEAL_DARK),
        ("Repeatability", "≥30 cycles\nwithin the agreed tolerance", BLUE_LIGHT, BLUE),
        ("Two interface levels", "User Mode for operation\nDeveloper Mode for tuning", ORANGE_LIGHT, ORANGE),
    ]
    for i, (title_value, body, fill, color) in enumerate(requirements):
        x = .58 + i * 3.78
        box(slide, x, 2.72, 3.43, 2.05, fill, color)
        text(slide, title_value, x + .24, 3.02, 2.93, .34, 16.5, True, NAVY)
        text(slide, body, x + .24, 3.56, 2.94, .78, 13.5, False, MID)

    box(slide, 11.92, 2.72, .78, 2.05, NAVY, None)
    text(slide, "PY", 12.02, 3.23, .58, .48, 21, True, WHITE, PP_ALIGN.CENTER)
    text(slide, "Python", 11.98, 3.85, .66, .25, 8.5, True, WHITE, PP_ALIGN.CENTER)

    box(slide, .58, 5.10, 12.12, 1.30, WHITE, LINE)
    text(slide, "WHY PYTHON", .85, 5.38, 1.48, .28, 10.5, True, TEAL_DARK)
    text(slide, "Fast UI iteration  •  scientific plotting  •  configurable models  •  repeatable tests  •  future hardware communication",
         2.34, 5.28, 9.92, .48, 15.5, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
    takeaway(slide, "Prototype scope today: one 250 mL lung branch; the software structure is designed to scale and be reused.")
    notes(slide, "Our control objective is to make every breath visible and adjustable. Users can set rate, tidal volume, I:E ratio, target flow and cycle count. We selected Python because it supports rapid interface development, scientific plotting, configuration files, repeatable simulations and later hardware communication. The current UI represents one 250 millilitre prototype branch, while the framework is designed for reuse.")
    return slide


def legend_item(slide, number, title_value, body, x, y):
    circle(slide, number, x, y, .34)
    text(slide, title_value, x + .47, y - .01, 2.60, .27, 12.5, True, NAVY)
    text(slide, body, x + .47, y + .29, 2.58, .55, 10.5, False, MID)


def user_mode_slide(prs):
    slide = blank(prs)
    content_header(slide, "User Mode: configure and review a breathing cycle")
    image = ASSETS / "ui_user_mode.png"
    slide.shapes.add_picture(str(image), Inches(.45), Inches(1.10),
                             width=Inches(8.70), height=Inches(5.22))
    for number, x, y in [(1, .68, 1.56), (2, .79, 4.06), (3, 4.31, 3.20),
                         (4, 8.28, 2.73), (5, .83, 5.01)]:
        circle(slide, number, x, y)
    box(slide, 9.38, 1.10, 3.46, 5.22, WHITE, LINE)
    legend_item(slide, 1, "Breathing settings", "Rate, volume, I:E, flow and cycle count.", 9.60, 1.36)
    legend_item(slide, 2, "Apply & playback", "Recalculate, start, pause, reset or stop.", 9.60, 2.31)
    legend_item(slide, 3, "Pneumatic state", "Route, valve opening, phase, flow and volume.", 9.60, 3.26)
    legend_item(slide, 4, "Live flow", "Simulated sensor flow versus requested setpoint.", 9.60, 4.21)
    legend_item(slide, 5, "Result", "Current frame and achieved breath versus target.", 9.60, 5.16)
    takeaway(slide, "User Mode keeps operation simple: choose the breath, apply it, and review the result.")
    notes(slide, "User Mode is designed for normal operation. The user enters the breathing target, applies it, and then plays the result. The centre shows the active atmosphere or vacuum route, pinch-valve opening, lung volume and flow. The chart compares flow with the requested setpoint, and the outcome panel reports the achieved breath. The current screen is simulation playback rather than live hardware feedback.")
    return slide


def developer_mode_slide(prs):
    slide = blank(prs)
    content_header(slide, "Developer Mode: tune, diagnose and calibrate")
    image = ASSETS / "ui_developer_mode.png"
    slide.shapes.add_picture(str(image), Inches(.45), Inches(1.10),
                             width=Inches(8.70), height=Inches(5.22))
    for number, x, y in [(1, .70, 1.66), (2, .70, 4.53), (3, 6.27, 1.40),
                         (4, 6.31, 3.06), (5, 6.31, 4.82)]:
        circle(slide, number, x, y)
    box(slide, 9.38, 1.10, 3.46, 5.22, WHITE, LINE)
    legend_item(slide, 1, "Test conditions", "Timing, target volume and number of cycles.", 9.60, 1.36)
    legend_item(slide, 2, "Control settings", "PID, feedforward, Smith predictor and guard.", 9.60, 2.31)
    legend_item(slide, 3, "Volume trace", "Target and repeated volume response.", 9.60, 3.26)
    legend_item(slide, 4, "Flow trace", "Tracking, overshoot and response shape.", 9.60, 4.21)
    legend_item(slide, 5, "Actuation & state", "PPV command, DAC output, route and timing.", 9.60, 5.16)
    takeaway(slide, "Developer Mode links performance to the controller, model assumptions, delays and valve commands.")
    notes(slide, "Developer Mode exposes the parameters needed for tuning and calibration. It aligns volume, flow, valve commands and the state machine on one timeline. The advanced settings include PID gains, model feedforward, Smith prediction, input and sensor delays, valve behaviour and lung-model parameters.")
    return slide


def model_slide(prs):
    slide = blank(prs)
    content_header(slide, "Control model: simple enough to tune, detailed enough to transfer")
    # Closed-loop flow diagram
    flow = [
        ("Breathing target", "RR · VT · I:E\nflow timing", TEAL_LIGHT, TEAL_DARK),
        ("State machine", "ATM / VAC routing\ndirection dead time", BLUE_LIGHT, BLUE),
        ("PID + model tools", "PI/PID · feedforward\nSmith prediction", ORANGE_LIGHT, ORANGE),
        ("Lung + sensor model", "capacity · time constants\ndelay · valve response", TEAL_LIGHT, TEAL_DARK),
    ]
    for i, (title_value, body, fill, color) in enumerate(flow):
        x = .54 + i * 2.67
        box(slide, x, 1.52, 2.35, 1.52, fill, color)
        text(slide, title_value, x + .18, 1.78, 1.98, .30, 14, True, NAVY,
             PP_ALIGN.CENTER)
        text(slide, body, x + .18, 2.20, 1.99, .53, 10.5, False, MID,
             PP_ALIGN.CENTER)
        if i < 3:
            text(slide, "→", x + 2.36, 2.03, .30, .35, 20, True, MID,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(slide, 11.20, 1.52, 1.58, 1.52, NAVY, None)
    text(slide, "Q(t)\nV(t)\ncommands", 11.40, 1.85, 1.18, .75, 13.5, True,
         WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    box(slide, .55, 3.46, 7.20, 2.82, WHITE, LINE)
    text(slide, "Adjustable model parameters", .84, 3.75, 3.10, .34, 17, True, NAVY)
    parameters = [
        "Lung capacity and inspiration/expiration response",
        "Pinch-valve nonlinearity and command delay",
        "Sensor delay, response time, noise and zero bias",
        "PID gains, feedforward and prediction settings",
    ]
    for i, item in enumerate(parameters):
        y = 4.30 + i * .44
        circle(slide, "✓", .87, y, .25, TEAL_DARK, 9)
        text(slide, item, 1.23, y - .02, 5.95, .28, 12.5, False, NAVY)

    box(slide, 8.05, 3.46, 4.73, 2.82, ORANGE_LIGHT, ORANGE)
    text(slide, "AUTO PID", 8.38, 3.73, 1.30, .34, 17, True, ORANGE)
    text(slide, "Offline candidate search", 9.62, 3.75, 2.63, .28, 12.5, True, NAVY)
    auto = [
        "Tests Kp / Ki candidates",
        "Scores 30-cycle results",
        "Checks error, consistency and saturation",
        "User reviews the best gains",
    ]
    for i, item in enumerate(auto):
        y = 4.30 + i * .44
        circle(slide, i + 1, 8.40, y, .25, ORANGE, 8.5)
        text(slide, item, 8.76, y - .02, 3.52, .28, 12.3, False, NAVY)
    takeaway(slide, "Auto PID supports tuning; physical measurements still define the model and safe operating limits.")
    notes(slide, "The control model follows the physical sequence: breathing targets define timing, the state machine selects atmosphere or vacuum, the controller sets pinch-valve opening, and the lung and sensor model produce flow and volume. The model parameters are adjustable rather than hard-coded. We also have an offline Auto-PID workflow that tests candidate gains over repeated breaths and scores volume error, consistency, flow tracking and saturation before the user applies a candidate.")
    return slide


def manual_slide(prs):
    slide = blank(prs)
    content_header(slide, "Built-in modeling & control guide")
    image = ASSETS / "ui_modeling_guide.png"
    slide.shapes.add_picture(str(image), Inches(.46), Inches(1.12),
                             width=Inches(8.62), height=Inches(5.45))
    box(slide, 9.31, 1.12, 3.52, 5.45, WHITE, LINE)
    text(slide, "The guide explains", 9.60, 1.44, 2.90, .34, 17, True, NAVY)
    items = [
        ("1", "System & setpoints", "Breathing sequence and timing equations"),
        ("2", "Physical parameters", "Meaning and effect of every model input"),
        ("3", "Delay model", "Command, transport and sensor response"),
        ("4", "PID, FF & Smith", "How the control options interact"),
        ("5", "Calibration", "Which tests identify each parameter"),
    ]
    for i, (num, title_value, body) in enumerate(items):
        y = 2.06 + i * .83
        circle(slide, num, 9.61, y, .32)
        text(slide, title_value, 10.05, y - .02, 2.38, .25, 12.2, True, NAVY)
        text(slide, body, 10.05, y + .25, 2.43, .39, 10.1, False, MID)
    takeaway(slide, "The manual connects each adjustable value to its physical meaning, equation and calibration method.")
    notes(slide, "Developer Mode includes a modeling and control guide. It is not only a parameter list: it explains the breathing sequence, timing equations, physical meaning of the model inputs, the delay chain, the PID, feedforward and Smith options, and the experiments needed to calibrate each parameter. This also documents how the complete system should be used.")
    return slide


def reusable_slide(prs):
    slide = blank(prs)
    content_header(slide, "Design philosophy: calibrate once, reuse the same framework")
    steps = [
        ("1", "Connect a prototype", "Different lung, tubing,\nvalves or sensor", TEAL_LIGHT, TEAL_DARK),
        ("2", "Characterise it", "Full-open response\nFixed PPV sweep", BLUE_LIGHT, BLUE),
        ("3", "Update parameters", "Capacity · τin · τout\nγ · delays · limits", ORANGE_LIGHT, ORANGE),
        ("4", "Validate & deploy", "30-cycle check\nApply reviewed gains", TEAL_LIGHT, TEAL_DARK),
    ]
    for i, (num, title_value, body, fill, color) in enumerate(steps):
        x = .48 + i * 3.18
        box(slide, x, 1.55, 2.74, 2.38, fill, color)
        circle(slide, num, x + .22, 1.82, .46, color, 14)
        text(slide, title_value, x + .22, 2.46, 2.30, .36, 15.5, True, NAVY,
             PP_ALIGN.CENTER)
        text(slide, body, x + .24, 3.02, 2.28, .58, 11.5, False, MID,
             PP_ALIGN.CENTER)
        if i < 3:
            text(slide, "→", x + 2.75, 2.55, .42, .40, 23, True, MID,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    box(slide, .72, 4.45, 5.75, 1.62, TEAL_LIGHT, TEAL_DARK)
    text(slide, "REUSED", 1.00, 4.73, 1.10, .29, 11, True, TEAL_DARK)
    text(slide, "UI workflow · state machine · data logging\nvalidation sequence · documentation structure",
         2.05, 4.65, 3.95, .70, 14, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
    box(slide, 6.84, 4.45, 5.75, 1.62, ORANGE_LIGHT, ORANGE)
    text(slide, "RE-IDENTIFIED", 7.13, 4.73, 1.48, .29, 11, True, ORANGE)
    text(slide, "Plant parameters · safe ranges · delays\ncontroller gains and hardware mappings",
         8.58, 4.65, 3.48, .70, 14, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
    takeaway(slide, "The prototype changes; the workflow stays. Measurements update the parameter file instead of rewriting the UI.")
    notes(slide, "Because both the lung and the pneumatic system are prototypes, we designed the software around a reusable workflow. For a different system, we run the same characterisation tests, update the fitted capacity, response, valve and delay parameters, then repeat validation and tuning. The interface, state machine, data logging and documentation remain the same. This process is also described in the built-in guide.")
    return slide


def bpm_slide(prs):
    slide = blank(prs)
    content_header(slide, "15 bpm achieved — now define the adjustable range")
    box(slide, .58, 1.34, 5.32, 4.98, TEAL_LIGHT, TEAL_DARK)
    circle(slide, "✓", 1.00, 1.76, .72, GREEN, 22)
    text(slide, "15.00 bpm", 1.95, 1.73, 2.96, .55, 27, True, NAVY)
    text(slide, "Completed validation", 1.04, 2.62, 2.50, .30, 13, True, TEAL_DARK)
    results = [
        ("30 cycles", "repeated simulation"),
        ("242.6–246.3 mL", "tidal-volume range"),
        ("Within ±5%", "of the 250 mL branch target"),
    ]
    for i, (value, label) in enumerate(results):
        y = 3.18 + i * .83
        text(slide, value, 1.05, y, 2.38, .32, 17, True, NAVY)
        text(slide, label, 3.25, y + .03, 2.08, .28, 11.5, False, MID)
    text(slide, "The nominal target is demonstrated in the current model benchmark.",
         1.05, 5.72, 4.28, .38, 11, False, TEAL_DARK)

    box(slide, 6.25, 1.34, 6.50, 4.98, ORANGE_LIGHT, ORANGE)
    text(slide, "INPUT NEEDED FROM 4DMEDICAL", 6.65, 1.72, 4.40, .30, 12, True, ORANGE)
    text(slide, "What adjustable respiratory-rate range should the final system support?",
         6.65, 2.20, 5.45, .76, 20, True, NAVY)
    questions = [
        "Minimum and maximum rate: 12–16, 12–18, or another range?",
        "Required adjustment step and accuracy?",
        "Must every rate hold the same 250 mL target and I:E = 1:2?",
    ]
    for i, item in enumerate(questions):
        y = 3.26 + i * .76
        circle(slide, i + 1, 6.67, y, .31, ORANGE, 9)
        text(slide, item, 7.12, y - .02, 4.92, .53, 14, True, NAVY)
    takeaway(slide, "Next requirement decision: define the operating envelope around the achieved 15 bpm nominal point.")
    notes(slide, "The nominal 15 breaths per minute target has been achieved in the current 30-cycle model benchmark, with a tidal-volume range of 242.6 to 246.3 millilitres, inside plus or minus five percent of the 250 millilitre branch target. The remaining question for 4DMedical is the required adjustable range: the minimum and maximum rate, adjustment step and whether the same volume and I:E ratio must hold across that full range.")
    return slide


def future_slide(prs):
    slide = blank(prs)
    content_header(slide, "Future work: migrate the model to the physical system")
    stages = [
        ("1", "Calibrate the path", "Full-open ATM/VAC response\nFixed-opening PPV sweep\nSynchronize Q(t), V(t), P(t),\ncommands", TEAL_LIGHT, TEAL_DARK, "Measured model"),
        ("2", "Close the real loop", "Drive PPV and routing valves\nFeed SFM3300 data to controller\nVerify timing and safety interlocks", BLUE_LIGHT, BLUE, "Hardware closed loop"),
        ("3", "Validate performance", "30 cycles at agreed\nnominal/boundary\nReport volume, flow and\nsaturation\nThen test independent dual\nbranches", ORANGE_LIGHT, ORANGE, "Requirement evidence"),
    ]
    for i, (num, title_value, body, fill, color, result) in enumerate(stages):
        x = .43 + i * 4.13
        box(slide, x, 1.20, 3.67, 5.03, fill, color)
        circle(slide, num, x + .28, 1.49, .50, color, 14)
        text(slide, title_value, x + .90, 1.51, 2.43, .55, 17, True, NAVY)
        line(slide, x + .28, 2.26, x + 3.37, 2.26, color, 1.4)
        text(slide, body, x + .30, 2.60, 3.00, 2.40, 13.2, False, NAVY)
        box(slide, x + .30, 5.44, 3.02, .49, WHITE, None, False)
        text(slide, result, x + .38, 5.55, 2.86, .21, 10.8, True, color,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < 2:
            text(slide, "→", x + 3.67, 3.19, .44, .42, 23, True, MID,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    takeaway(slide, "Success is a measured 30-cycle result against the operating range agreed with 4DMedical.")
    notes(slide, "Migration to the physical system has three stages. First, calibrate the assembled atmosphere and vacuum paths and the pinch valve using synchronized flow, independent volume, pressure and command data. Second, close the real hardware loop so the controller drives the valves and uses SFM3300 feedback. Third, validate 30 cycles at the agreed nominal and boundary conditions before extending to two independent branches.")
    return slide


def team_architecture_slide(prs, picture_blobs):
    """Create a stable replacement for the unsupported SmartArt-like slide."""
    if len(picture_blobs) != 3:
        raise RuntimeError("Expected three architecture pictures on original slide 3")
    slide = blank(prs)
    content_header(slide, "Team responsibilities & system architecture")
    roles = [
        ("01", "Mechanical Part", "Chun Yin Ku (John)",
         "Pneumatic architecture · aerodynamic sizing\nphysical interface", TEAL_LIGHT, TEAL_DARK),
        ("02", "Electrical Part", "Jiaquan",
         "Power distribution · PCB design\nsensor interfacing", BLUE_LIGHT, BLUE),
        ("03", "Control Software & UI", "Yuqiao",
         "STM32 control · flow-profile shaping\nreal-time UI", ORANGE_LIGHT, ORANGE),
    ]
    for i, (num, title_value, owner, body, fill, color) in enumerate(roles):
        y = 1.20 + i * 1.77
        box(slide, .48, y, 6.28, 1.50, fill, color)
        circle(slide, num, .75, y + .28, .53, color, 14)
        text(slide, title_value, 1.52, y + .22, 2.75, .34, 17, True, NAVY)
        text(slide, owner, 4.25, y + .24, 2.04, .28, 11.5, True, color,
             PP_ALIGN.RIGHT)
        text(slide, body, 1.52, y + .76, 4.78, .54, 12.5, False, MID)

    box(slide, 7.05, 1.20, 5.80, 5.04, WHITE, LINE)
    text(slide, "Integrated system", 7.35, 1.43, 2.50, .31, 15.5, True, NAVY)
    slide.shapes.add_picture(BytesIO(picture_blobs[0]), Inches(7.38), Inches(1.88),
                             width=Inches(2.28), height=Inches(2.52))
    slide.shapes.add_picture(BytesIO(picture_blobs[2]), Inches(10.28), Inches(1.82),
                             width=Inches(1.46), height=Inches(2.46))
    text(slide, "+", 9.72, 2.73, .42, .40, 22, True, MID,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    line(slide, 7.34, 4.54, 12.56, 4.54, LINE, 1)
    text(slide, "Control interface", 7.35, 4.67, 2.50, .26, 11.5, True, TEAL_DARK)
    slide.shapes.add_picture(BytesIO(picture_blobs[1]), Inches(7.38), Inches(5.01),
                             width=Inches(5.12), height=Inches(1.84))
    notes(slide, "The project is divided into three connected responsibilities. John covers the mechanical and pneumatic design, Jiaquan covers the electrical power and sensor interfaces, and Yuqiao covers STM32 control, flow shaping and the real-time Python interface. The right side shows how the board, pneumatic architecture and software interface form one system.")
    return slide


def delete_after(prs, keep_count):
    slide_ids = prs.slides._sldIdLst
    while len(slide_ids) > keep_count:
        slide_id = slide_ids[-1]
        rel_id = slide_id.rId
        slide_ids.remove(slide_id)
        prs.part.drop_rel(rel_id)


def delete_slide(prs, index):
    slide_ids = prs.slides._sldIdLst
    slide_id = slide_ids[index]
    rel_id = slide_id.rId
    slide_ids.remove(slide_id)
    prs.part.drop_rel(rel_id)


def build():
    required = [
        SOURCE,
        ASSETS / "ui_user_mode.png",
        ASSETS / "ui_developer_mode.png",
        ASSETS / "ui_modeling_guide.png",
    ]
    for path in required:
        if not path.exists():
            raise FileNotFoundError(path)

    prs = Presentation(SOURCE)
    delete_after(prs, 11)
    picture_blobs = [
        shape.image.blob for shape in prs.slides[2].shapes
        if getattr(shape, "shape_type", None) == 13
    ]
    original_ids = list(prs.slides._sldIdLst)

    team_architecture = team_architecture_slide(prs, picture_blobs)
    agenda = agenda_slide(prs)
    mechanical = transition(
        prs, "01", "Mechanical Part",
        "Pneumatic architecture  •  airflow sizing  •  physical prototype",
    )
    notes(mechanical, "Part one presents the mechanical design, including the pneumatic architecture, airflow sizing, prototype connections, airtightness testing and assembly plan.")
    electrical = transition(
        prs, "02", "Electrical Part",
        "Power distribution  •  bridge board  •  sensor and valve signals",
    )
    notes(electrical, "Part two presents the electrical control architecture, bridge board and signal verification.")
    control = transition(
        prs, "03", "Control Software & UI",
        "Python interface  •  adjustable model  •  Auto PID  •  reusable workflow",
    )
    notes(control, "Part three presents the Python control software and interface, including the user and developer modes, adjustable model, Auto-PID workflow and physical-system migration plan.")

    control_slides = [
        objective_slide(prs),
        user_mode_slide(prs),
        developer_mode_slide(prs),
        model_slide(prs),
        manual_slide(prs),
        reusable_slide(prs),
        bpm_slide(prs),
        future_slide(prs),
    ]

    new_ids = list(prs.slides._sldIdLst)[len(original_ids):]
    team_architecture_id, agenda_id, mechanical_id, electrical_id, control_id = new_ids[:5]
    control_ids = new_ids[5:]
    assert len(control_ids) == len(control_slides)

    # Delete the old unsupported SmartArt slide only after all replacement
    # slides have unique package part names.
    delete_slide(prs, 2)

    desired = (
        [original_ids[0], agenda_id, mechanical_id]
        + [original_ids[1], team_architecture_id]
        + original_ids[3:8]
        + [electrical_id]
        + original_ids[8:11]
        + [control_id]
        + control_ids
    )
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        slide_id_list.remove(slide_id)
    for slide_id in desired:
        slide_id_list.append(slide_id)

    prs.core_properties.title = "LungLogic — Mechanical, Electrical and Control UI"
    prs.core_properties.subject = "Integrated capstone presentation for 4DMedical"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
