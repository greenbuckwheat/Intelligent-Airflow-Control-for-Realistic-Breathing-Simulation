"""Append an editable, five-minute control-simulation section to the team deck."""

from pathlib import Path
from lxml import etree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\13922\Downloads\23-09-Present.pptx")
OUT = Path(__file__).with_name("23-09-Present_ControlSimulation.pptx")
SCRIPT = Path(__file__).with_name("five_minute_speaker_notes.md")
EXPERIMENT_SCRIPT = Path(__file__).with_name("experiment_design_speaker_notes.md")
INK = RGBColor(16, 36, 63)
BODY = RGBColor(26, 32, 44)
TEAL = RGBColor(14, 128, 127)
BLUE = RGBColor(33, 105, 186)
RED = RGBColor(186, 64, 55)
PALE = RGBColor(246, 250, 255)
MINT = RGBColor(231, 247, 246)
CREAM = RGBColor(255, 243, 224)
GREY = RGBColor(99, 112, 126)


def box(slide, x, y, w, h, fill=PALE, border=None, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if border is None else None
    if border is not None:
        shape.line.color.rgb = border
    return shape


def txt(slide, value, x, y, w, h, size=16, bold=False, color=BODY,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(4)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return shape


def line(slide, x1, y1, x2, y2, color=GREY, width=1.5):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def slide(prs, title, number, source, takeaway, note):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    txt(s, title, .38, .22, 11.0, .53, 24, True, INK)
    txt(s, "4DMedical  /  LungLogic", 11.75, .25, 1.2, .32, 6.5, True, TEAL,
        PP_ALIGN.RIGHT)
    line(s, .38, .84, 12.95, .84, RGBColor(221, 231, 238), 1)
    box(s, .38, 6.72, 12.57, .43, MINT)
    txt(s, takeaway, .52, 6.77, 12.2, .31, 11, True, INK)
    txt(s, source, .4, 7.20, 11.9, .19, 6.5, False, GREY)
    txt(s, str(number), 12.35, 7.18, .55, .19, 7, False, GREY, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = note
    return s


def card(s, x, y, w, h, heading, body, fill=PALE, heading_color=INK,
         body_size=14):
    box(s, x, y, w, h, fill)
    txt(s, heading, x + .17, y + .16, w - .34, .42, 17, True, heading_color)
    txt(s, body, x + .17, y + .68, w - .34, h - .78, body_size,
        valign=MSO_ANCHOR.TOP)


def fix_team_slide_route(prs):
    """Correct teammate slide 5 in the submitted copy; keep the source intact."""
    root = prs.slides[4]._element
    blocks = [e for e in root.iter() if etree.QName(e).localname == "AlternateContent"]
    for block in blocks:
        runs = [e for e in block.iter() if etree.QName(e).localname == "t"]
        texts = [e.text or "" for e in runs]
        if any("Vacuum solenoid opens" in t for t in texts):
            for e in runs:
                if "Vacuum solenoid opens" in (e.text or ""):
                    e.text = " Atmosphere solenoid opens; the lung expands and takes in air."
                elif "Atmospheric solenoid opens" in (e.text or ""):
                    e.text = " Vacuum solenoid opens; regulated suction contracts the lung."
                elif (e.text or "").strip() in {"0", "kPa gauge."}:
                    e.text = ""


def add_control_slide(prs):
    s = slide(prs, "Python control simulation: one branch, one cycle", 11,
              "Source: single_lung_pid_gui.py; PROTOTYPE_REQUIREMENTS.md",
              "Atmosphere opens for inspiration; vacuum opens for expiration.",
              "I will now move from the hardware architecture to the Python control model. "
              "One branch represents a 250 millilitre share of the two-lung target. "
              "At 15 breaths per minute the period is four seconds. With an inspiration "
              "to expiration ratio of one to two, inspiration is 1.33 seconds and expiration "
              "is 2.67 seconds. A 100 hertz state machine selects atmosphere or vacuum, "
              "while the pinch valve sets the opening. The model then estimates flow, "
              "volume and a delayed sensor reading. The graph in the UI is playback of a "
              "precomputed simulation, not live hardware closed-loop data.")
    card(s, .5, 1.10, 3.0, 3.25, "User settings",
         "15 bpm  |  I:E = 1:2\n250 mL target per branch\n4.00 s cycle\n1.33 s in / 2.67 s out", MINT)
    card(s, 3.72, 1.10, 3.0, 3.25, "100 Hz controller",
         "State machine + phase timing\nFlow PI + model feedforward\nOptional Smith prediction\nVolume deadline guard", PALE)
    card(s, 6.94, 1.10, 2.75, 3.25, "Commands",
         "Atmosphere → lung expands\nVacuum → lung contracts\nPinch valve (PPV) meters flow\n80 ms direction dead time", CREAM)
    card(s, 9.91, 1.10, 2.9, 3.25, "Simulated plant",
         "Lung volume and flow\nValve/transport lag\nDelayed, noisy flow sensor\nIntegrated-volume estimate", PALE)
    for x in [3.52, 6.74, 9.71]:
        txt(s, "→", x, 2.48, .2, .42, 19, True, TEAL)
    txt(s, "Architecture is implemented in Python; STM32 demo reads SFM3300 but does not feed this PID.",
        .63, 4.76, 12.1, .65, 15, True, INK)
    return s


def parameter_table(s, rows):
    columns = [("Parameter / current value", .55, 2.57),
               ("Physical meaning", 3.16, 3.28),
               ("Evidence now → how to establish it", 6.48, 6.30)]
    for label, x, w in columns:
        box(s, x, 1.10, w, .42, INK, rounded=False)
        txt(s, label, x + .10, 1.13, w - .20, .32, 11, True,
            RGBColor(255, 255, 255))
    row_h = 4.90 / len(rows)
    for i, (value, meaning, evidence) in enumerate(rows):
        y = 1.57 + i * row_h
        fill = PALE if i % 2 == 0 else RGBColor(255, 255, 255)
        for _, x, w in columns:
            box(s, x, y, w, row_h - .03, fill, rounded=False)
        txt(s, value, .68, y + .05, 2.36, row_h - .11, 11, True, INK,
            valign=MSO_ANCHOR.TOP)
        txt(s, meaning, 3.27, y + .05, 3.05, row_h - .11, 10.7, color=BODY,
            valign=MSO_ANCHOR.TOP)
        txt(s, evidence, 6.59, y + .05, 6.06, row_h - .11, 10.7, color=BODY,
            valign=MSO_ANCHOR.TOP)


def add_parameters_slide_a(prs):
    s = slide(prs, "Model parameters I: timing, capacity and rebound", 12,
              "Source: PROTOTYPE_REQUIREMENTS.md; Orifice.mlx; single_lung_config.gui.json",
              "Only the breathing target and software sample rate are set; plant values still need calibration.",
              "These rows separate what we set from what we know about the real plant. "
              "Fifteen breaths per minute, a one-to-two ratio and 250 millilitres per "
              "branch are design targets, so they are verified by measuring actual cycles "
              "and delivered volume. One hundred hertz is a software setting; command "
              "timestamps will verify the real timing. The 250 millilitre capacity is "
              "configured, and needs a measured chamber volume. The 1.30 second fully "
              "open endpoint has one visual observation but is not a repeatable complete-path "
              "curve. The 0.57 second inspiration constant comes from an orifice/rebound "
              "candidate. We must measure and fit volume and flow throughout a full breath.")
    parameter_table(s, [
        ("15 bpm; I:E 1:2; 250 mL", "Requested cycle and tidal volume; 4.00 s total, 1.33 s in, 2.67 s out.",
         "Design target. Measure full-path cycles and delivered volume to see whether it is achieved."),
        ("Control loop 100 Hz", "Model and command update every 10 ms.",
         "Software setting. Timestamp real commands to check actual period and jitter."),
        ("Capacity Vmax 250 mL", "Upper bound on modelled change in lung volume.",
         "Nominal container capacity. Verify by calibrated displacement/geometry and pressure-to-full volume; compare with integrated flow."),
        ("Full-open fill 1.30 s", "From collapsed lung to visually full with atmosphere path fully open.",
         "One visual observation only. Repeat synchronized full-path volume–time and flow–time traces; inspect start/end criteria."),
        ("Inspiration τin 0.57 s", "Shape of atmosphere-driven filling: early or late volume gain.",
         "Candidate from Orifice.mlx, not a full-path fit. Record entire fill curve and calibrate with measured capacity and flow integral."),
        ("Model scope", "Combined inflation/deflation curves, valve and sensor response.",
         "No separate sponge compliance, tubing resistance, lung pressure or changing vacuum pressure; measure these before claiming physical parameters."),
    ])
    return s


def add_parameters_slide_b(prs):
    s = slide(prs, "Model parameters II: valve, delays and controller", 13,
              "Source: PROTOTYPE_REQUIREMENTS.md; single_lung_config.gui.json; single_lung_pid_gui.py",
              "Measure valve and sensor dynamics before trusting simulated flow tracking or PID tuning.",
              "The rest of the table describes the pathway from a valve command to a sensor "
              "reading. The expiration time constant is assumed because suction, sponge "
              "rebound and tubing restriction have not been separately measured. The pinch "
              "valve exponent approximates nonlinear flow versus opening, and needs a voltage "
              "sweep with checks for hysteresis. The command-to-lung and lung-to-sensor delays "
              "are assumed twenty-five milliseconds each; the sensor time constant is twenty "
              "milliseconds. We need synchronized command and flow traces to separate them. "
              "Noise, zero offset and eighty millisecond direction dead time also need "
              "measurement. Kp, Ki, Kd and Smith delay are controller settings, not physical "
              "lung parameters. Finally, compare thirty-cycle volume error and oscillation "
              "after the plant is calibrated.")
    parameter_table(s, [
        ("Expiration τout 0.50 s", "How quickly vacuum contracts the lung; combines lung response and path restriction.",
         "Assumed. At the intended vacuum pressure, record fully open volume–time and flow–time traces and fit."),
        ("PPV exponent γ 1.15", "Nonlinear relation between commanded pinch-valve opening and actual flow.",
         "Assumed. Sweep DAC voltage/opening at fixed pressure and initial volume; record flow and check hysteresis."),
        ("Command → lung 25 ms", "Delay from control command to flow response at the lung.",
         "Assumed. Timestamp command, valve electrical signal and lung-end flow separately; 100 Hz rounds delay to steps."),
        ("Lung → sensor 25 ms; τs 20 ms", "Transport delay and sensor rise/fall speed after flow changes.",
         "Assumed. Record simultaneous lung-end reference flow and sensor output; align onsets and fit response."),
        ("Noise 0.10 SLM; bias 0", "Random reading fluctuation and steady zero-flow offset.",
         "Assumed. Sample closed-valve zero flow repeatedly; calculate mean and variation."),
        ("Direction dead time 80 ms", "Interval with both routes closed during atmosphere/vacuum switching.",
         "Controller policy. Log both valve commands and actual flow to verify order and effective switching time."),
        ("Kp .01; Ki .06; Kd 0; Smith 45 ms", "Controller corrects flow error and estimates effects of delayed feedback.",
         "Tuning settings, not lung physics. Calibrate valve/sensor first; compare 30-cycle volume error, oscillation and saturation."),
    ])
    return s


def add_benchmark_slide(prs):
    s = slide(prs, "Completed test: 30 simulated breathing cycles", 14,
              "Source: output/requirement_benchmark/single_lung_30_cycle_benchmark.md",
              "The present assumed plant passes at 15 bpm but is at its limit by 16 bpm.",
              "This chart shows the minimum and maximum tidal volume over 30 simulated "
              "cycles, with the volume deadline guard enabled and internal true volume used "
              "for feedback. The requirement band starts at 237.5 millilitres, five percent "
              "below the 250 millilitre target. Twelve and fifteen breaths per minute pass. "
              "At sixteen, the minimum is 237.3, just 0.2 millilitres below the band; eighteen "
              "fails clearly. With only flat-flow PI and no deadline guard at fifteen, the "
              "range is about 215 to 220 millilitres. This is a simulated sensitivity test, "
              "not a measurement of achievable hardware performance.")
    box(s, .58, 1.12, 8.13, 4.92, PALE)
    x0, x1, lo, hi = 2.23, 8.02, 210.0, 252.0
    fx = lambda v: x0 + (v - lo) / (hi - lo) * (x1 - x0)
    for v in [210, 220, 230, 240, 250]:
        x = fx(v)
        line(s, x, 1.62, x, 5.58, RGBColor(222, 231, 238), .8)
        txt(s, str(v), x - .23, 5.62, .48, .29, 10, color=GREY, align=PP_ALIGN.CENTER)
    line(s, fx(237.5), 1.57, fx(237.5), 5.54, RED, 1.8)
    txt(s, "−5% limit: 237.5 mL", 4.3, 1.23, 2.3, .26, 10, True, RED)
    rows = [("12 bpm", 245.3, 250.0, "PASS"),
            ("15 bpm", 242.6, 246.3, "PASS"),
            ("16 bpm", 237.3, 240.8, "FAIL"),
            ("18 bpm", 227.2, 231.1, "FAIL")]
    for i, (label, low, high, status) in enumerate(rows):
        y = 2.02 + i * .87
        txt(s, label, .8, y - .17, 1.14, .33, 13, True)
        line(s, fx(low), y, fx(high), y, TEAL if status == "PASS" else RED, 7)
        txt(s, f"{low:.1f}–{high:.1f}", 6.96, y - .16, 1.02, .32, 12, True,
            TEAL if status == "PASS" else RED, PP_ALIGN.RIGHT)
    txt(s, "Tidal volume / mL", 3.56, 5.73, 3.1, .25, 10, color=GREY,
        align=PP_ALIGN.CENTER)
    card(s, 8.96, 1.12, 3.77, 2.08, "Model-only result",
         "The 15 bpm case meets ±5% for 30 simulated cycles.", MINT, body_size=14)
    card(s, 8.96, 3.4, 3.77, 2.64, "Control limitation",
         "Without the volume guard, 15 bpm gives 214.9–219.8 mL. At 16–18 bpm the assumed full-open plant becomes limiting.",
         CREAM, body_size=13)
    return s


def add_tuning_slide(prs):
    s = slide(prs, "Auto PI tuning: paper-inspired offline study", 15,
              "Sources: Killingsworth & Krstić, ACC 2005; single_lung_pid_es_study.py; output/pid_es_study/report.md",
              "Gain tuning improves the simulated PI result slightly; the flow-capacity limit remains.",
              "The cited paper studies online PID auto-tuning through extremum seeking. "
              "Our implementation is a simpler offline, bounded two-sided perturbation "
              "search over proportional and integral gains; it is inspired by the paper, "
              "not an exact reproduction of its algorithm. It runs twenty iterations, "
              "scoring thirty simulated cycles at fifteen breaths per minute. Derivative "
              "gain remains zero. Without the volume guard, the tidal range improves from "
              "212.8 to 217.8 millilitres to 215.2 to 220.2, but both miss the five percent "
              "requirement. With the guard, both reach about 243 to 246 millilitres, yet the "
              "pinch valve is fully open for almost all inspiration. Controller gains cannot "
              "make a saturated pneumatic path fill faster.")
    card(s, .58, 1.12, 3.85, 4.75, "Method",
         "Offline two-sided gain search\n20 iterations, 30 cycles each\nKp and Ki tuned; Kd = 0\n15 bpm, guard OFF", PALE, body_size=16)
    card(s, 4.66, 1.12, 3.85, 4.75, "Numerical outcome",
         "Kp: 0.010 → 0.01198\nKi: 0.060 → 0.240\nBaseline: 212.8–217.8 mL\nTuned: 215.2–220.2 mL\nBoth fail ±5%", MINT, body_size=16)
    card(s, 8.74, 1.12, 4.0, 4.75, "What this tells us",
         "Guard ON: ≈243–246 mL\nPPV open ≈100% in inspiration\nActuator/plant capacity dominates\nPhysical plant still unmeasured", CREAM, body_size=16)
    return s


def add_ui_slide(prs):
    s = slide(prs, "User and Developer modes: controls and checks", 16,
              "Source: single_lung_pid_gui.py; single_lung_config.gui.json",
              "The interface separates operating choices from uncalibrated model and control settings.",
              "The interface now has a User mode and a Developer mode. User mode exposes "
              "breathing choices such as rate, target volume and inspiration to expiration "
              "ratio, and presents the simulated curves and status. Developer mode exposes "
              "the model's time constants, capacity, delays, PID gains, Smith predictor and "
              "volume guard. Invalid values such as impossible phase timing, fill time or "
              "flow target are checked and reported through a GUI error dialog. There is "
              "also a diagnostic warning when volume looks close to target only because "
              "the pinch valve is fully open. These checks prevent misleading simulation "
              "settings; they do not prove a hardware safety system. Displayed values are "
              "simulated, and the STM32 flow reading is not yet used by this model.")
    card(s, .58, 1.12, 3.86, 4.74, "User mode",
         "Respiratory rate\nTidal-volume target\nI:E ratio / flow setting\nSimulation cycle count\nPlayback and status display", MINT, body_size=16)
    card(s, 4.71, 1.12, 3.86, 4.74, "Developer mode",
         "Capacity and time constants\nValve curve and delays\nPI / Smith parameters\nSensor and volume guard", PALE, body_size=17)
    card(s, 8.84, 1.12, 3.89, 4.74, "Input & diagnostic feedback",
         "Invalid configuration → error dialog\nTiming / target feasibility checks\nWarnings for PPV saturation\nNo live hardware feedback yet", CREAM, body_size=16)
    return s


def add_next_slide(prs):
    s = slide(prs, "Next: measure the real pneumatic path", 17,
              "Source: PROTOTYPE_REQUIREMENTS.md; current simulation and sensitivity studies",
              "Use measured flow–time curves to refit the model, then rerun the 30-cycle tests.",
              "The next step is to instrument the assembled pneumatic path, then run two "
              "calibration experiments: a fully open inspiration and expiration response, "
              "and a fixed-opening pinch-valve sweep. Both require synchronized flow, volume, "
              "pressure and command data. If a firmer sponge is installed, we will repeat "
              "the same fully open response under matched conditions and compare its curves. "
              "We will use these data to refit the model before repeating thirty-cycle and "
              "two-branch tests. The next slides specify the experiments; none of the "
              "measurements shown there have already been performed.")
    card(s, .58, 1.12, 5.95, 4.85, "1. Record & calibrate",
         "Synchronize flow(t), volume(t), pressure(t)\nLog ATM/VAC commands and PPV voltage\nSweep fixed PPV openings in both directions\nFit endpoint, τ, delays and sensor response", MINT, body_size=16)
    card(s, 6.78, 1.12, 5.95, 4.85, "2. Refit & verify",
         "Repeat full-open response if sponge changes\nRefit the model with assembled-path data\nRepeat 30-cycle tests across 12–18 bpm\nTest two-branch operation and consistency", CREAM, body_size=16)
    return s


def add_experiment_overview_slide(prs):
    s = slide(prs, "Two experiments to identify the real plant", 18,
              "Source: PROTOTYPE_REQUIREMENTS.md; current parameter evidence table",
              "Both experiments use the complete assembled path; the curves become calibration data.",
              "We propose exactly two calibration experiments. First, open the pinch valve "
              "fully and measure the entire atmosphere-inspiration and vacuum-expiration "
              "responses. This determines the actual fill endpoint, the two curve shapes, "
              "peak flow and residual volume. Any future sponge replacement is compared by "
              "repeating this same experiment under the same conditions. Second, hold the "
              "pinch valve at several fixed commands in each flow direction. This identifies "
              "the relationship between command and flow, including dead zone and hysteresis. "
              "Pressure and starting volume must be controlled or recorded because they also "
              "change the flow. These are proposed measurements, not results we already have.")
    card(s, .60, 1.16, 5.90, 4.94, "1. Fully open response",
         "What: ATM inspiration and VAC expiration with PPV fully open.\n\n"
         "Purpose: measure the complete path's flow and volume response, not just an endpoint.\n\n"
         "Outputs: Q(t), V(t), P(t), time to 250 mL, τin, τout, peak flow, residual volume and repeatability.",
         MINT, body_size=15)
    card(s, 6.76, 1.16, 5.96, 4.94, "2. Fixed PPV opening sweep",
         "What: hold several PPV commands, separately for ATM and VAC routes.\n\n"
         "Purpose: establish command-to-flow behavior at matched pressure and starting volume.\n\n"
         "Outputs: flow versus command, effective γ, dead zone, hysteresis, saturation and command-to-flow delay.",
         CREAM, body_size=15)
    return s


def add_experiment_setup_slide(prs):
    s = slide(prs, "How to run experiment 1: full-open response", 19,
              "Sources: README_STM32_DEMO_CN.md; PROTOTYPE_REQUIREMENTS.md; SFM3300 datasheet",
              "Synchronized flow, pressure, volume and command traces identify the fill limit and delays.",
              "Once the whole atmosphere and vacuum path is assembled, a proposed Python "
              "experiment program will select test conditions and save time-stamped data. "
              "A microcontroller or DAQ should execute the timed valve sequence and acquire "
              "signals; the present STM32 demo cannot yet command the pinch valve or provide "
              "physical PID feedback. Measure flow with SFM3300, pressure at the source and "
              "lung, and lung volume independently, for example by calibrated displacement. "
              "First confirm zero flow and sign. Start from the same collapsed volume; open "
              "the ATM route and fully open PPV until full, then open VAC to empty. Repeat "
              "at the intended source pressures. Find time to 250 millilitres from measured "
              "volume, fit inspiration and expiration curves, and compare integrated flow "
              "against independent volume. Log command and actual electrical transitions "
              "to estimate delays. A new sponge can be assessed by repeating this protocol.")
    card(s, .55, 1.10, 3.88, 4.95, "Build & instrument",
         "Assemble final lung + tubing + PPV + ATM/VAC path.\n\n"
         "Measure SFM3300 flow, source/lung pressures and independent lung volume.\n\n"
         "Record valve drive/GPIO and PPV voltage; use one time base.", PALE, body_size=14)
    card(s, 4.69, 1.10, 3.88, 4.95, "Run a repeatable step",
         "Confirm zero/sign; start from the same collapsed volume.\n\n"
         "At fixed source pressure: ATM + PPV 100% until full; VAC + PPV 100% until empty.\n\n"
         "Repeat runs. If sponge changes, repeat identical steps.", MINT, body_size=14)
    card(s, 8.83, 1.10, 3.92, 4.95, "Calculate from CSV",
         "Vref(t): time to 250 mL, final/residual volume, repeatability.\n\n"
         "Q(t), P(t): peak flow, fill/empty curves; fit τin and τout.\n\n"
         "Command/electrical/flow onset: total delay; components need a fast lung-end reference.", CREAM, body_size=14)
    return s


def add_experiment_sweep_slide(prs):
    s = slide(prs, "How to run experiment 2: PPV opening sweep", 20,
              "Sources: single_lung_pid_gui.py; README_STM32_DEMO_CN.md",
              "Fit the opening-to-flow relationship only after matching pressure and lung state.",
              "The same acquisition setup can run the second experiment. The proposed Python "
              "program schedules a sequence of fixed PPV commands, while the MCU or DAQ "
              "applies them and records time-stamped signals. For each command, start from a "
              "matched lung volume and source pressure, hold the ATM or VAC route fixed, then "
              "collect the initial flow response and a defined measurement window. Test both "
              "rising and falling commands, with repeats, to reveal dead zone and hysteresis. "
              "Fit the model exponent from flow versus opening under matched conditions; "
              "compare early and late flow because sponge volume changes during each run. "
              "Timestamp the valve voltage and the first flow response to estimate total "
              "command-to-flow delay. One flow sensor alone cannot separate valve motion, "
              "gas transport and sensor lag, so that decomposition needs electrical signals "
              "and an independent lung-end reference flow. The current GUI/firmware do not "
              "yet perform this experiment automatically.")
    card(s, .56, 1.10, 3.82, 4.97, "Command sequence",
         "Choose multiple fixed DAC / PPV settings from closed to fully open.\n\n"
         "Repeat each point for ATM and VAC, at matched starting volume and pressure.\n\n"
         "Sweep up, then down, with replicate runs.", PALE, body_size=14)
    card(s, 4.63, 1.10, 3.82, 4.97, "Log synchronized data",
         "Python saves one CSV per run: timestamps, route, PPV command and measured voltage.\n\n"
         "Also record Q(t), Psource(t), Plung(t) and independent Vref(t).\n\n"
         "MCU/DAQ performs timed switching and sampling.", MINT, body_size=14)
    card(s, 8.70, 1.10, 4.06, 4.97, "Identify parameters",
         "At matched pressure and volume, plot initial flow versus PPV setting.\n\n"
         "Fit effective γ; report dead zone, hysteresis and saturation.\n\n"
         "Measure command-to-flow delay; extra reference sensing separates its causes.", CREAM, body_size=14)
    return s


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    prs = Presentation(SOURCE)
    if len(prs.slides) != 10:
        raise RuntimeError(f"Expected 10 team slides, found {len(prs.slides)}")
    fix_team_slide_route(prs)
    for fn in (add_control_slide, add_parameters_slide_a, add_parameters_slide_b,
               add_benchmark_slide, add_tuning_slide, add_ui_slide, add_next_slide,
               add_experiment_overview_slide, add_experiment_setup_slide,
               add_experiment_sweep_slide):
        fn(prs)
    prs.save(OUT)
    headings = [
        "11. Python control simulation (45 s)",
        "12. Model parameters I: timing and rebound (45 s)",
        "13. Model parameters II: valve and sensor (55 s)",
        "14. Completed 30-cycle study (45 s)",
        "15. Offline auto PI study (55 s)",
        "16. User/Developer modes and errors (40 s)",
        "17. Next experiments and limitation (50 s)",
    ]
    parts = ["# Five-minute continuation: Python control simulation\n",
             "English script for slides 11–17. Estimated duration: about 5–6 minutes "
             "at a measured speaking pace. The same text is embedded in slide notes.\n"]
    for heading, s in zip(headings, list(prs.slides)[10:]):
        parts.extend([f"## {heading}\n", s.notes_slide.notes_text_frame.text.strip() + "\n"])
    parts.append(
        "## 中文备讲要点\n\n"
        "- 气路方向：大气通路对应吸气、肺体积增加；真空通路对应呼气、肺体积减少。\n"
        "- 250 mL 是单支路目标；1.30 s 满充终点及 τ 等尚不是整机实测参数。\n"
        "- 30 周期与自动 PI 结果均为离线仿真；自动调参尚未接入实时硬件或 GUI 操作按钮。\n"
        "- 16 bpm 基准仿真最低 237.3 mL，比 5% 下限低 0.2 mL；不要声称硬件结论。\n"
        "- 参数页的三列分别是当前值、物理意义、当前证据和验证实验；按行解释，不要只念数值。\n"
        "- 需测完整气路流量-时间曲线，并尝试较硬海绵、验证回弹是否真的加快。\n"
    )
    SCRIPT.write_text("\n".join(parts), encoding="utf-8")
    experiment_headings = [
        "18. Two experiments: purpose and outputs",
        "19. Fully open response: setup, sequence and calculations",
        "20. PPV sweep: setup, sequence and calculations",
    ]
    experiment_parts = [
        "# Experiment-design addendum (slides 18–20)\n",
        "These are planned measurements. The existing STM32 demo only reads the SFM3300 "
        "and displays simulated commands; it does not drive PPV or use physical feedback.\n",
    ]
    for heading, s in zip(experiment_headings, list(prs.slides)[17:]):
        experiment_parts.extend([f"## {heading}\n",
                                 s.notes_slide.notes_text_frame.text.strip() + "\n"])
    experiment_parts.append(
        "## 中文备讲要点\n\n"
        "- 只计划两个实验：全开吸气／呼气响应，固定 PPV 开度扫描。更换海绵后重复第一个实验，不另列对比实验。\n"
        "- 完整回路接好后，Python 负责设定方案和保存数据；MCU／DAQ 负责有确定时间戳的阀门执行与传感器采样。当前演示固件还没有这些功能。\n"
        "- 每个试验至少同步保存 Q(t)、独立 Vref(t)、气源及肺端 P(t)、PPV 命令／实测电压、ATM/VAC 命令及实际电信号。\n"
        "- 仅有一个流量计只能估计总命令到读数延迟；要拆分阀门、气路与传感器延迟，需电信号和快速肺端参考流量或体积传感器。\n"
    )
    EXPERIMENT_SCRIPT.write_text("\n".join(experiment_parts), encoding="utf-8")
    print(f"Saved {OUT} ({len(prs.slides)} slides)")
    print(f"Saved {SCRIPT}")
    print(f"Saved {EXPERIMENT_SCRIPT}")


if __name__ == "__main__":
    main()
