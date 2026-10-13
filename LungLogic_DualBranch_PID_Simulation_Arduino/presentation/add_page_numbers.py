"""Add unobtrusive slide numbers without altering existing slide content."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


SOURCE = Path(r"C:\Users\13922\Desktop\capstone\23-09-Presentation_Integrated_Control_UI.pptx")
OUTPUT = SOURCE.with_name(SOURCE.stem + "_Numbered.pptx")

# Title slide and the three section-opening slides intentionally have no number.
SKIP_SLIDES = {1, 4, 10, 14}


def add_page_number(slide, number: int) -> None:
    shape = slide.shapes.add_textbox(
        Inches(12.20), Inches(7.08), Inches(0.70), Inches(0.26)
    )
    shape.name = f"Page Number {number}"
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = str(number)
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(12)
    paragraph.font.color.rgb = RGBColor(91, 106, 119)


def main() -> None:
    prs = Presentation(SOURCE)
    for number, slide in enumerate(prs.slides, start=1):
        if number not in SKIP_SLIDES:
            add_page_number(slide, number)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
