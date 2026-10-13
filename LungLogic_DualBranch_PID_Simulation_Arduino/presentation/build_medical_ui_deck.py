"""Build a concise, editable UI-focused presentation for 4DMedical."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "LungLogic_UI_Medical_Group_Presentation.pptx"

NAVY = RGBColor(20, 43, 66)
NAVY_2 = RGBColor(32, 66, 92)
TEAL = RGBColor(0, 137, 145)
TEAL_LIGHT = RGBColor(226, 245, 245)
BLUE = RGBColor(48, 111, 171)
BLUE_LIGHT = RGBColor(232, 241, 249)
ORANGE = RGBColor(229, 126, 49)
ORANGE_LIGHT = RGBColor(255, 241, 226)
RED = RGBColor(192, 68, 58)
GREEN = RGBColor(34, 137, 97)
INK = RGBColor(31, 42, 51)
MID = RGBColor(91, 105, 117)
LIGHT = RGBColor(246, 248, 250)
LINE = RGBColor(216, 224, 230)
WHITE = RGBColor(255, 255, 255)

W = 13.333
H = 7.5


def add_text(slide, text, x, y, w, h, size=18, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(3)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return shape


def add_box(slide, x, y, w, h, fill=LIGHT, border=None, radius=True):
    kind = (MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius
            else MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(1)
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.2, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dash is not None:
        shape.line.dash_style = dash
    return shape


def add_circle(slide, number, x, y, diameter=.36, fill=TEAL):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = WHITE
    shape.line.width = Pt(1.5)
    add_text(slide, str(number), x, y + .005, diameter, diameter - .01, 13, True,
             WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return shape


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def header(slide, title, section, number, takeaway=None):
    add_text(slide, section.upper(), .52, .22, 3.6, .22, 8.5, True, TEAL)
    add_text(slide, title, .50, .50, 12.2, .55, 25, True, NAVY)
    add_line(slide, .5, 1.12, 12.83, 1.12, LINE, 1)
    if takeaway:
        add_box(slide, .5, 6.72, 12.33, .38, TEAL_LIGHT, None, False)
        add_text(slide, takeaway, .66, 6.78, 11.85, .22, 10.5, True, NAVY,
                 valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, str(number), 12.55, 7.18, .28, .15, 7.5, False, MID,
             PP_ALIGN.RIGHT)


def source(slide, text):
    add_text(slide, text, .52, 7.18, 11.6, .15, 7, False, MID)


def base_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def title_slide(prs):
    s = base_slide(prs)
    add_box(s, 0, 0, W, H, NAVY, None, False)
    add_box(s, .55, .58, .16, 5.95, TEAL, None, False)
    add_text(s, "LungLogic UI", 1.02, 1.15, 7.6, .8, 34, True, WHITE)
    add_text(s, "From breathing requirements to testable control", 1.02, 2.04,
             10.9, .65, 24, False, RGBColor(204, 230, 233))
    add_text(s, "Capstone update for 4DMedical / Medical Group", 1.04, 3.03,
             8.7, .4, 16, False, WHITE)
    add_box(s, 1.02, 4.23, 10.95, 1.2, NAVY_2, RGBColor(77, 115, 140))
    add_text(s, "Focus: what the interface does, what the simulation shows,\nand which system requirements still need agreement",
             1.32, 4.49, 10.2, .7, 17, True, WHITE, valign=MSO_ANCHOR.MIDDLE)
    add_text(s, "22 September 2026", 1.03, 6.76, 4.0, .22, 10, False,
             RGBColor(190, 207, 220))
    set_notes(s, "Today I will focus on the interface: how it turns breathing requirements into a repeatable simulation, what has been implemented, and where the physical system currently limits the result. I will finish with the decisions and measurements needed from the next stage.")


def objective_slide(prs):
    s = base_slide(prs)
    header(s, "Objective first: make breathing performance visible and testable",
           "1 — Objective & requirements", 2,
           "Prototype scope: one 250 mL branch; dual-branch performance is not yet demonstrated.")
    add_box(s, .62, 1.46, 12.05, 1.18, TEAL_LIGHT, TEAL)
    add_text(s, "OBJECTIVE", .91, 1.68, 1.55, .25, 10, True, TEAL)
    add_text(s, "Let a user configure a breathing target, see the pneumatic response, and judge whether repeated breaths meet the target before hardware deployment.",
             2.37, 1.60, 9.9, .68, 18, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
    requirements = [
        ("01", "250 mL per branch", "500 mL total system target"),
        ("02", "12–18 bpm", "15 bpm nominal; I:E = 1:2"),
        ("03", "30 repeatable cycles", "tidal volume within ±5%"),
        ("04", "Independent branches", "safe ATM / VAC switching"),
    ]
    for i, (num, title, detail) in enumerate(requirements):
        x = .62 + i * 3.05
        add_box(s, x, 3.02, 2.78, 2.35, WHITE, LINE)
        add_text(s, num, x + .18, 3.22, .48, .35, 13, True, TEAL)
        add_text(s, title, x + .18, 3.78, 2.42, .42, 17, True, NAVY)
        add_text(s, detail, x + .18, 4.37, 2.40, .57, 12.5, False, MID)
    source(s, "Requirements source: project brief and PROTOTYPE_REQUIREMENTS.md")
    set_notes(s, "The objective is not simply to display a moving lung. The interface converts a requested breath into timing, flow and valve behaviour that can be checked. The system-level targets are 500 millilitres across two branches, 12 to 18 breaths per minute with 15 nominal, a one-to-two I:E ratio, and repeatability over 30 breaths. This prototype represents only one 250 millilitre branch, so it cannot yet claim the two-branch requirement.")


def divider(prs, number, title, subtitle, slide_no):
    s = base_slide(prs)
    add_box(s, 0, 0, W, H, NAVY, None, False)
    add_text(s, f"{number:02d}", .83, 1.15, 1.2, .9, 34, True, TEAL)
    add_text(s, title, 2.23, 1.28, 9.6, .8, 31, True, WHITE)
    add_line(s, 2.23, 2.36, 11.9, 2.36, RGBColor(75, 104, 126), 1.4)
    add_text(s, subtitle, 2.23, 2.74, 9.5, 1.25, 20, False,
             RGBColor(204, 220, 231), valign=MSO_ANCHOR.MIDDLE)
    add_text(s, str(slide_no), 12.55, 7.18, .28, .15, 7.5, False,
             RGBColor(160, 182, 198), PP_ALIGN.RIGHT)
    return s


def legend_item(slide, number, title, body, x, y):
    add_circle(slide, number, x, y, .34)
    add_text(slide, title, x + .47, y - .01, 2.35, .27, 12.5, True, NAVY)
    add_text(slide, body, x + .47, y + .29, 2.35, .58, 10.5, False, MID)


def user_mode_slide(prs):
    s = base_slide(prs)
    header(s, "User Mode: set a breath, play it back, assess the result",
           "2 — UI walkthrough", 4,
           "The interface exposes clinical-style settings while keeping model assumptions out of the main workflow.")
    img = ASSETS / "ui_user_mode.png"
    s.shapes.add_picture(str(img), Inches(.50), Inches(1.40), width=Inches(8.65), height=Inches(5.19))
    # Figure annotations
    for n, x, y in [(1, .73, 1.85), (2, .84, 4.32), (3, 4.34, 3.50),
                    (4, 8.26, 3.05), (5, .88, 5.28)]:
        add_circle(s, n, x, y)
    add_box(s, 9.38, 1.40, 3.45, 5.19, WHITE, LINE)
    legend_item(s, 1, "Breathing inputs", "Rate, tidal volume, I:E, flow and cycle count.", 9.60, 1.66)
    legend_item(s, 2, "Apply & playback", "Recalculate, then start, pause, reset or stop.", 9.60, 2.62)
    legend_item(s, 3, "Pneumatic state", "ATM/VAC route, PPV opening, phase, flow and volume.", 9.60, 3.58)
    legend_item(s, 4, "Live flow", "Measured model output against the flat setpoint.", 9.60, 4.54)
    legend_item(s, 5, "Outcome", "Current frame plus last-breath volume versus target.", 9.60, 5.50)
    source(s, "Figure: current single_lung_pid_gui.py User Mode; displayed values are simulated")
    set_notes(s, "User Mode is the main demonstration view. First, the user enters the breathing target. Applying the settings recalculates the full simulation. Playback then shows the current phase, routing valves, pinch-valve opening, sensor location, lung volume and flow. The panel on the left reports the current frame and the most recent tidal volume. The flow chart compares simulated measured flow with the requested flat setpoint. This is precomputed playback, not live hardware data.")


def developer_mode_slide(prs):
    s = base_slide(prs)
    header(s, "Developer Mode: expose assumptions and diagnose why a case passes",
           "2 — UI walkthrough", 5,
           "A passing volume is interpreted together with flow tracking, saturation, timing and valve state.")
    img = ASSETS / "ui_developer_mode.png"
    s.shapes.add_picture(str(img), Inches(.50), Inches(1.40), width=Inches(8.65), height=Inches(5.19))
    for n, x, y in [(1, .76, 1.96), (2, .76, 4.81), (3, 6.29, 1.69),
                    (4, 6.33, 3.33), (5, 6.33, 5.08)]:
        add_circle(s, n, x, y)
    add_box(s, 9.38, 1.40, 3.45, 5.19, WHITE, LINE)
    legend_item(s, 1, "Test definition", "Breathing timing, target volume and run length.", 9.60, 1.66)
    legend_item(s, 2, "Controller options", "PID, feedforward, Smith prediction and volume guard.", 9.60, 2.62)
    legend_item(s, 3, "Volume evidence", "True model volume and target over repeated breaths.", 9.60, 3.58)
    legend_item(s, 4, "Flow evidence", "Sensor-model flow reveals overshoot and missed setpoints.", 9.60, 4.54)
    legend_item(s, 5, "Actuation & state", "PPV opening, DAC command, route and phase timing.", 9.60, 5.50)
    source(s, "Figure: current single_lung_pid_gui.py Developer Mode; displayed values are simulated")
    set_notes(s, "Developer Mode answers a different question: why did the result pass or fail? It exposes timing, controller gains, model feedforward, Smith prediction, input delay, sensor assumptions and the optional volume deadline guard. The plots align volume, flow, PPV command and the state machine on the same time axis. This matters because a tidal-volume pass may still rely on valve saturation or a large flow pulse.")


def status_capability_slide(prs):
    s = base_slide(prs)
    header(s, "The UI is implemented; the physical closed loop is not yet verified",
           "3 — Current status", 7,
           "Use the UI today for simulation, requirement discussion and experiment planning—not hardware acceptance.")
    add_box(s, .62, 1.48, 5.86, 4.92, TEAL_LIGHT, TEAL)
    add_text(s, "AVAILABLE NOW", .94, 1.77, 2.3, .31, 12, True, TEAL)
    implemented = [
        "User and Developer workflows",
        "Input checks and timing warnings",
        "Flow, volume, command and state plots",
        "Repeatable 30-cycle simulation studies",
        "Saved configuration and playback",
    ]
    for i, item in enumerate(implemented):
        y = 2.34 + i * .69
        add_circle(s, "✓", .94, y, .31, GREEN)
        add_text(s, item, 1.40, y + .01, 4.63, .34, 16, i == 3, NAVY)

    add_box(s, 6.82, 1.48, 5.86, 4.92, ORANGE_LIGHT, ORANGE)
    add_text(s, "NOT YET DEMONSTRATED", 7.14, 1.77, 3.2, .31, 12, True, ORANGE)
    pending = [
        "Physical PPV actuation from this controller",
        "SFM3300 feedback used by the PID",
        "Synchronized pressure acquisition",
        "30-cycle hardware repeatability",
        "Independent two-branch operation",
    ]
    for i, item in enumerate(pending):
        y = 2.34 + i * .69
        add_circle(s, "–", 7.14, y, .31, ORANGE)
        add_text(s, item, 7.60, y + .01, 4.67, .34, 16, False, NAVY)
    source(s, "Evidence: source code, STM32 demo scope, PROTOTYPE_REQUIREMENTS.md")
    set_notes(s, "The interface itself is implemented and supports two levels of use, parameter validation, plots, repeatable simulations and saved settings. However, the physical controller is not complete. The STM32 demo can read the SFM3300, but that reading is not yet used by this PID model; the pinch valve is not driven by this loop; pressure is not synchronized; and neither 30 physical breaths nor independent dual-branch control has been verified.")


def benchmark_slide(prs):
    s = base_slide(prs)
    header(s, "At 15 bpm, thirty simulated breaths pass with almost no margin",
           "3 — Current status", 8,
           "This is a model result under stated assumptions, not evidence of installed-path performance.")
    # Plot frame and scale
    add_box(s, .62, 1.46, 8.38, 4.92, WHITE, LINE)
    x0, x1 = 2.25, 8.42
    lo, hi = 220.0, 252.0
    fx = lambda value: x0 + (value - lo) / (hi - lo) * (x1 - x0)
    for value in (220, 230, 237.5, 240, 250):
        x = fx(value)
        add_line(s, x, 1.94, x, 5.78, RED if value == 237.5 else LINE,
                 1.8 if value == 237.5 else .8)
        add_text(s, f"{value:g}", x - .30, 5.86, .60, .25, 9.5,
                 value == 237.5, RED if value == 237.5 else MID, PP_ALIGN.CENTER)
    add_text(s, "±5% lower limit", fx(237.5) - .70, 1.63, 1.42, .23, 9.5,
             True, RED, PP_ALIGN.CENTER)
    rows = [
        ("12 bpm", 245.3, 250.0, True),
        ("15 bpm", 242.6, 246.3, True),
        ("16 bpm", 237.3, 240.8, False),
        ("18 bpm", 227.2, 231.1, False),
    ]
    for i, (label, low, high, passed) in enumerate(rows):
        y = 2.35 + i * .82
        color = GREEN if passed else RED
        add_text(s, label, .90, y - .18, 1.05, .32, 13, True, NAVY)
        add_line(s, fx(low), y, fx(high), y, color, 8)
        add_circle(s, "✓" if passed else "×", 8.52, y - .16, .32, color)
        add_text(s, f"{low:.1f}–{high:.1f} mL", 6.62, y - .18, 1.65, .32,
                 11.5, True, color, PP_ALIGN.RIGHT)
    add_text(s, "Tidal volume across 30 cycles / mL", 3.72, 6.10, 3.35, .25,
             10, False, MID, PP_ALIGN.CENTER)

    add_box(s, 9.28, 1.46, 3.40, 2.26, TEAL_LIGHT, TEAL)
    add_text(s, "15 bpm", 9.60, 1.78, 1.1, .38, 20, True, NAVY)
    add_text(s, "242.6–246.3 mL\nPasses the ±5% simulation band", 9.60, 2.27,
             2.72, .85, 14, True, GREEN)
    add_box(s, 9.28, 3.92, 3.40, 2.46, ORANGE_LIGHT, ORANGE)
    add_text(s, "Critical caveat", 9.60, 4.22, 2.45, .32, 14, True, ORANGE)
    add_text(s, "PPV is fully open for ~100% of inspiration.\nNo meaningful throttling margin remains.",
             9.60, 4.76, 2.69, 1.12, 14, True, NAVY)
    source(s, "Data: output/requirement_benchmark/single_lung_30_cycle_benchmark.md; guard on, assumed plant")
    set_notes(s, "The 30-cycle model result is useful, but it must be described precisely. Twelve and fifteen breaths per minute are within the 237.5 to 262.5 millilitre acceptance band. Sixteen misses by 0.2 millilitres in its weakest breath, and eighteen fails clearly. At fifteen, however, the volume deadline guard drives the pinch valve fully open for almost every inspiration sample. Therefore fifteen is a marginal simulated pass with no useful control margin, not a robust hardware result.")


def physical_limit_slide(prs):
    s = base_slide(prs)
    header(s, "15 bpm is a requirement trade-off, not just a PID tuning problem",
           "3 — Current status", 9,
           "When the PPV is already fully open, software cannot create additional pneumatic flow.")
    add_text(s, "Available inspiration time at 15 bpm, I:E = 1:2", .66, 1.44,
             6.7, .34, 14, True, NAVY)
    # Timeline
    add_box(s, .68, 2.08, 7.20, .70, BLUE_LIGHT, BLUE, False)
    add_box(s, .68, 2.08, .44, .70, ORANGE, None, False)
    add_box(s, 1.12, 2.08, .14, .70, RED, None, False)
    add_text(s, "80 ms\nswitch", .66, 2.91, .62, .48, 10, True, ORANGE,
             PP_ALIGN.CENTER)
    add_text(s, "~25 ms\ndelay", 1.17, 2.91, .72, .48, 10, True, RED,
             PP_ALIGN.CENTER)
    add_text(s, "≈1.23 s remaining for pneumatic response", 2.00, 2.28,
             4.93, .30, 17, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(s, "1.333 s inspiration slot", .68, 1.80, 7.20, .22, 10.5,
             False, MID, PP_ALIGN.CENTER)
    add_text(s, "Assumed full-open reach: 1.30 s", .68, 3.61, 7.20, .33,
             18, True, RED, PP_ALIGN.CENTER)
    add_line(s, 1.26, 3.45, 7.88, 3.45, RED, 3)
    add_text(s, "Result: the controller is saturated before it can create comfortable margin.",
             .90, 4.27, 6.78, .68, 17, True, NAVY, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE)
    add_box(s, 8.26, 1.46, 4.42, 4.92, ORANGE_LIGHT, ORANGE)
    add_text(s, "DECISION NEEDED FROM 4DMEDICAL", 8.58, 1.79, 3.83, .33,
             11.5, True, ORANGE)
    prompts = [
        "Must 15 bpm, 250 mL and I:E 1:2 all be fixed simultaneously?",
        "Is ±5% required for every breath, and is a high early-flow pulse acceptable?",
        "If the path saturates, should we change timing, target volume or hardware capacity?",
    ]
    for i, prompt in enumerate(prompts):
        y = 2.39 + i * 1.10
        add_circle(s, i + 1, 8.61, y, .34, ORANGE)
        add_text(s, prompt, 9.10, y - .04, 3.06, .78, 14.5, True, NAVY)
    add_text(s, "18 bpm under the current assumed plant: 227–231 mL; PID gain changes cannot overcome full-open saturation.",
             8.59, 5.73, 3.58, .45, 10.5, False, RED)
    source(s, "Timing and saturation evidence: benchmark and 18 bpm trade-off studies; installed path remains unmeasured")
    set_notes(s, "At fifteen breaths per minute, the full cycle is four seconds and the one-to-two ratio gives 1.333 seconds for inspiration. Direction switching and the assumed input delay leave about 1.23 seconds of pneumatic response, while the model assumes 1.30 seconds to reach full volume from empty. That explains why the volume guard saturates the valve. We need 4DMedical to tell us which constraints are clinically or experimentally non-negotiable: cadence, tidal volume, I:E ratio, per-breath tolerance and flow shape. At eighteen, this assumed plant cannot reach the tolerance band even when fully open, so changing PID gains alone is not a solution.")


def future_slide(prs):
    s = base_slide(prs)
    header(s, "Future work: replace assumptions with measurements, then validate",
           "4 — Future", 11,
           "Success is a measured 30-cycle result against an agreed requirement—not a favourable simulation trace.")
    stages = [
        ("1", "Calibrate the path", "Full-open ATM/VAC response\nFixed-opening PPV sweep\nSynchronize Q(t), V(t), P(t), commands", TEAL_LIGHT, TEAL),
        ("2", "Close the real loop", "Drive PPV and routing valves\nFeed SFM3300 data to controller\nVerify timing and safety interlocks", BLUE_LIGHT, BLUE),
        ("3", "Validate performance", "30 cycles at agreed nominal/boundary\nReport volume, flow and saturation\nThen test independent dual branches", ORANGE_LIGHT, ORANGE),
    ]
    for i, (num, title, body, fill, color) in enumerate(stages):
        x = .63 + i * 4.17
        add_box(s, x, 1.58, 3.72, 4.33, fill, color)
        add_circle(s, num, x + .27, 1.88, .52, color)
        add_text(s, title, x + .93, 1.91, 2.47, .42, 18, True, NAVY)
        add_line(s, x + .28, 2.57, x + 3.43, 2.57, color, 1.2)
        add_text(s, body, x + .31, 2.92, 3.05, 1.70, 14.5, False, NAVY)
        label = ["Measured model", "Hardware closed loop", "Requirement evidence"][i]
        add_box(s, x + .31, 5.04, 3.05, .48, WHITE, None, False)
        add_text(s, label, x + .39, 5.13, 2.89, .22, 11.5, True, color,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < 2:
            add_text(s, "→", x + 3.72, 3.38, .45, .42, 24, True, MID,
                     PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    source(s, "Proposed experiments: experiment_design_speaker_notes.md and PROTOTYPE_REQUIREMENTS.md")
    set_notes(s, "The next stage has three gates. First, calibrate the assembled path using a fully open inspiration and expiration test plus a fixed-opening pinch-valve sweep, with synchronized flow, independent volume, pressure and command signals. Second, implement the real closed loop so the microcontroller drives the valves and uses measured flow. Third, run 30-cycle tests at the agreed nominal and boundary conditions, then extend to two independently controlled branches. Each gate produces evidence for the next one.")


def build():
    for image in (ASSETS / "ui_user_mode.png", ASSETS / "ui_developer_mode.png"):
        if not image.exists():
            raise FileNotFoundError(f"Missing {image}; run capture_ui_for_medical_deck.py")
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    title_slide(prs)
    objective_slide(prs)
    s = divider(prs, 2, "UI walkthrough",
                "One simple mode for setting and reviewing a breath; one technical mode for understanding the controller and assumptions.", 3)
    set_notes(s, "I will now show the two interface layers. User Mode is designed around the requested breath and its outcome. Developer Mode opens the controller and model details needed for diagnosis and calibration.")
    user_mode_slide(prs)
    developer_mode_slide(prs)
    s = divider(prs, 3, "Current status",
                "What is implemented, what the 30-cycle study shows, and why the nominal 15 bpm case still needs a requirements decision.", 6)
    set_notes(s, "The next section separates completed software capability from physical-system evidence, then examines the nominal fifteen-breath-per-minute case.")
    status_capability_slide(prs)
    benchmark_slide(prs)
    physical_limit_slide(prs)
    s = divider(prs, 4, "Future",
                "Calibrate the installed pneumatic path, close the hardware loop, and validate against the agreed acceptance test.", 10)
    set_notes(s, "The final section turns the current uncertainty into a short measurement and validation plan.")
    future_slide(prs)
    prs.core_properties.title = "LungLogic UI — Medical Group Presentation"
    prs.core_properties.subject = "Objective, requirements, UI functions, current status and future work"
    prs.core_properties.author = "LungLogic Capstone Team"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
